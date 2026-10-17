"""
Batch Document Generation Service.
Processes multiple validated document jobs in isolation, supports partial success,
and packages generated artifacts into a clean downloadable ZIP archive.
"""
import uuid
import zipfile
import shutil
from pathlib import Path
from typing import Dict, Any, List
from app.utils.file_manager import file_manager, sanitize_filename
from app.utils import supabase_storage
from app.schemas.executive_meet import ExecutiveMeetRequest
from app.schemas.bill_certificate import BillCertificateRequest
from app.schemas.bill_summary import BillSummaryRequest
from app.schemas.prize_money_receipt import PrizeMoneyReceiptRequest
from app.services.executive_meet_service import ExecutiveMeetService
from app.services.bill_certificate_service import BillCertificateService
from app.services.bill_summary_service import BillSummaryService
from app.services.prize_money_receipt_service import PrizeMoneyReceiptService
from app.services.document_repository import DocumentRepository
import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

BATCH_DIR = Path(__file__).resolve().parent.parent.parent / "generated_docs" / "batches"
BATCH_DIR.mkdir(parents=True, exist_ok=True)

class BatchService:
    MAX_BATCH_SIZE = 15

    @classmethod
    def process_batch(cls, jobs: List[Dict[str, Any]], user_id: str = "batch-user") -> Dict[str, Any]:
        if not jobs:
            raise ValueError("Batch request must contain at least one job.")
        if len(jobs) > cls.MAX_BATCH_SIZE:
            raise ValueError(f"Batch size exceeds maximum limit of {cls.MAX_BATCH_SIZE} jobs.")

        batch_id = f"batch-{uuid.uuid4().hex[:12]}"
        batch_folder = BATCH_DIR / batch_id
        batch_folder.mkdir(parents=True, exist_ok=True)

        job_results = []
        successful_files = []
        successful_count = 0
        failed_count = 0

        for idx, job in enumerate(jobs):
            job_id = job.get("job_id") or f"job-{idx+1}"
            doc_type = job.get("document_type")
            payload = job.get("payload", {})

            result_entry = {
                "job_id": job_id,
                "document_type": doc_type,
                "status": "pending",
                "error": None,
                "files": {}
            }

            try:
                if doc_type == "exec_meet":
                    validated = ExecutiveMeetRequest(**payload)
                    safe_date = sanitize_filename(validated.event_date.replace("/", "-"))
                    docx_fn = f"{job_id}_Executive_Meet_{safe_date}.docx"
                    pdf_fn = f"{job_id}_Executive_Meet_{safe_date}.pdf"
                    
                    docx_p = batch_folder / docx_fn
                    pdf_p = batch_folder / pdf_fn
                    
                    ExecutiveMeetService.generate_docx(validated, docx_p)
                    ExecutiveMeetService.generate_pdf(validated, pdf_p)
                    
                    # Also persist in repository
                    DocumentRepository.create_document(
                        user_id=user_id,
                        document_type="exec_meet",
                        title=validated.event_name,
                        data=validated.model_dump(),
                        docx_path=str(docx_p),
                        pdf_path=str(pdf_p),
                        change_summary=f"Batch generated ({job_id})"
                    )

                elif doc_type == "bill_certificate":
                    validated = BillCertificateRequest(**payload)
                    safe_date = sanitize_filename(validated.date.replace("/", "-"))
                    docx_fn = f"{job_id}_Bill_Certificate_{safe_date}.docx"
                    pdf_fn = f"{job_id}_Bill_Certificate_{safe_date}.pdf"
                    
                    docx_p = batch_folder / docx_fn
                    pdf_p = batch_folder / pdf_fn
                    
                    BillCertificateService.generate_docx(validated, docx_p)
                    BillCertificateService.generate_pdf(validated, pdf_p)
                    
                    DocumentRepository.create_document(
                        user_id=user_id,
                        document_type="bill_certificate",
                        title=f"Bill Certificate - {validated.chapter_name}",
                        data=validated.model_dump(),
                        docx_path=str(docx_p),
                        pdf_path=str(pdf_p),
                        change_summary=f"Batch generated ({job_id})"
                    )

                elif doc_type == "bill_summary":
                    validated = BillSummaryRequest(**payload)
                    safe_date = sanitize_filename(validated.date.replace("/", "-"))
                    docx_fn = f"{job_id}_Bill_Summary_{safe_date}.docx"
                    pdf_fn = f"{job_id}_Bill_Summary_{safe_date}.pdf"
                    
                    docx_p = batch_folder / docx_fn
                    pdf_p = batch_folder / pdf_fn
                    
                    BillSummaryService.generate_docx(validated, docx_p)
                    BillSummaryService.generate_pdf(validated, pdf_p)
                    
                    DocumentRepository.create_document(
                        user_id=user_id,
                        document_type="bill_summary",
                        title=validated.title,
                        data=validated.model_dump(),
                        docx_path=str(docx_p),
                        pdf_path=str(pdf_p),
                        change_summary=f"Batch generated ({job_id})"
                    )
                elif doc_type == "prize_money_receipt":
                    validated = PrizeMoneyReceiptRequest(**payload)
                    safe_date = sanitize_filename(validated.date.replace("/", "-"))
                    docx_fn = f"{job_id}_Prize_Money_Receipt_{safe_date}.docx"
                    pdf_fn = f"{job_id}_Prize_Money_Receipt_{safe_date}.pdf"

                    docx_p = batch_folder / docx_fn
                    pdf_p = batch_folder / pdf_fn

                    PrizeMoneyReceiptService.generate_docx(validated, docx_p)
                    PrizeMoneyReceiptService.generate_pdf(validated, pdf_p)

                    DocumentRepository.create_document(
                        user_id=user_id,
                        document_type="prize_money_receipt",
                        title=f"Prize Receipt - {validated.student_name}",
                        data=validated.model_dump(mode="json"),
                        docx_path=str(docx_p),
                        pdf_path=str(pdf_p),
                        change_summary=f"Batch generated ({job_id})"
                    )
                elif doc_type == "submitted":
                    submission_id = payload.get("submission_id", f"SUB-{uuid.uuid4().hex[:6].upper()}")
                    title = payload.get("title", "Official Submission Manifest & Dossier")
                    date_str = payload.get("date", "15/10/2026")
                    submitted_by = payload.get("submitted_by", "ACM Student Chapter")
                    department = payload.get("department", "Computer Science & Engineering, SVNIT")
                    remarks = payload.get("remarks", "Document submitted for official workflow verification.")

                    safe_date = sanitize_filename(date_str.replace("/", "-"))
                    docx_fn = f"{job_id}_Submitted_{safe_date}.docx"
                    pdf_fn = f"{job_id}_Submitted_{safe_date}.pdf"

                    docx_p = batch_folder / docx_fn
                    pdf_p = batch_folder / pdf_fn

                    # Generate DOCX
                    doc = docx.Document()
                    doc.add_heading("SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT", level=1)
                    doc.add_heading("ASSOCIATION FOR COMPUTING MACHINERY (ACM)", level=2)
                    doc.add_heading(f"SUBMISSION MANIFEST: {title}", level=3)
                    p = doc.add_paragraph()
                    p.add_run(f"Submission ID: {submission_id}\n").bold = True
                    p.add_run(f"Date: {date_str}\n")
                    p.add_run(f"Submitted By: {submitted_by}\n")
                    p.add_run(f"Department: {department}\n")
                    p.add_run(f"Remarks: {remarks}\n")
                    p.add_run("Status: SUBMITTED (Pending Faculty Endorsement)\n")
                    doc.save(docx_p)

                    # Generate PDF
                    doc_pdf = SimpleDocTemplate(str(pdf_p), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
                    styles = getSampleStyleSheet()
                    story = [
                        Paragraph("<b>SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT</b>", styles["Title"]),
                        Paragraph("<b>ACM STUDENT CHAPTER — SUBMISSION MANIFEST</b>", styles["Heading2"]),
                        Spacer(1, 14),
                        Paragraph(f"<b>Dossier Title:</b> {title}", styles["Normal"]),
                        Paragraph(f"<b>Submission ID:</b> {submission_id}", styles["Normal"]),
                        Paragraph(f"<b>Date:</b> {date_str}", styles["Normal"]),
                        Paragraph(f"<b>Submitted By:</b> {submitted_by}", styles["Normal"]),
                        Paragraph(f"<b>Department:</b> {department}", styles["Normal"]),
                        Spacer(1, 10),
                        Paragraph(f"<b>Remarks / Purpose:</b> {remarks}", styles["Normal"]),
                        Spacer(1, 10),
                        Paragraph("<b>Current Status:</b> SUBMITTED &amp; REGISTERED IN WORKFLOW", styles["Normal"]),
                        Spacer(1, 20),
                        Paragraph("_______________________<br/>Authorized Submitter Signature", styles["Normal"])
                    ]
                    doc_pdf.build(story)

                    DocumentRepository.create_document(
                        user_id=user_id,
                        document_type="submitted",
                        title=title,
                        data=payload,
                        docx_path=str(docx_p),
                        pdf_path=str(pdf_p),
                        change_summary=f"Batch generated ({job_id})"
                    )
                else:
                    raise ValueError(f"Unsupported document type: '{doc_type}'.")

                result_entry["status"] = "completed"
                result_entry["files"] = {
                    "docx": docx_fn,
                    "pdf": pdf_fn
                }
                successful_files.extend([docx_p, pdf_p])
                successful_count += 1

            except Exception as e:
                result_entry["status"] = "failed"
                result_entry["error"] = str(e)
                failed_count += 1

            job_results.append(result_entry)

        # Create ZIP archive if there are successful files
        zip_path = None
        if successful_files:
            zip_filename = f"Batch_{batch_id}.zip"
            zip_path = batch_folder / zip_filename
            with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
                for file_p in successful_files:
                    zf.write(file_p, arcname=file_p.name)
            if supabase_storage.storage_backend() == "supabase":
                supabase_storage.upload_file(
                    f"batches/{batch_id}/{zip_filename}", zip_path, "application/zip"
                )

        return {
            "batch_id": batch_id,
            "total_jobs": len(jobs),
            "successful_count": successful_count,
            "failed_count": failed_count,
            "jobs": job_results,
            "zip_filename": f"Batch_{batch_id}.zip" if zip_path else None,
            "zip_download_url": f"/api/documents/batch/{batch_id}/zip" if zip_path else None
        }

    @classmethod
    def get_batch_zip(cls, batch_id: str) -> Path:
        if not batch_id or any(ch not in "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_" for ch in batch_id):
            raise FileNotFoundError(f"Batch archive for '{batch_id}' not found.")
        batch_folder = BATCH_DIR / batch_id
        zip_path = batch_folder / f"Batch_{batch_id}.zip"
        if supabase_storage.storage_backend() == "supabase":
            from app.utils.file_manager import file_manager
            reference = supabase_storage.storage_reference(
                f"batches/{batch_id}/Batch_{batch_id}.zip", "application/zip"
            )
            try:
                materialized, _ = file_manager.materialize(reference)
                return materialized
            except RuntimeError as exc:
                raise FileNotFoundError(f"Batch archive for '{batch_id}' not found.") from exc
        if not zip_path.exists():
            raise FileNotFoundError(f"Batch archive for '{batch_id}' not found.")
        return zip_path
