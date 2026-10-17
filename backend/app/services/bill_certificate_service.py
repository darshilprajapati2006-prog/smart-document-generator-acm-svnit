from pathlib import Path
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether
)

from app.core.config import settings
from app.schemas.bill_certificate import BillCertificateRequest

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_table_borders(table, border_color="94A3B8"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

class BillCertificateService:
    @staticmethod
    def calculate_totals(data: BillCertificateRequest):
        """Perform reliable server-side calculations."""
        calculated_items = []
        subtotal = 0.0
        for idx, item in enumerate(data.items, start=1):
            total_cost = round(float(item.quantity) * float(item.unit_cost), 2)
            subtotal += total_cost
            calculated_items.append({
                "sr_no": idx,
                "name": item.name,
                "quantity": item.quantity,
                "unit_cost": float(item.unit_cost),
                "total_cost": total_cost,
            })
        subtotal = round(subtotal, 2)
        taxes = round(float(data.taxes_or_other_charges), 2)
        grand_total = round(subtotal + taxes, 2)
        return calculated_items, subtotal, taxes, grand_total

    @classmethod
    def generate_docx(cls, data: BillCertificateRequest, output_path: Path):
        calc_items, subtotal, taxes, grand_total = cls.calculate_totals(data)
        doc = Document()
        
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.55)
            section.right_margin = Inches(0.55)

        # Header with Logo & Institute Details
        header_table = doc.add_table(rows=1, cols=3)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False
        widths = [Inches(1.0), Inches(5.0), Inches(1.2)]
        for i, w in enumerate(widths):
            header_table.rows[0].cells[i].width = w

        # SVNIT Logo left
        svnit_logo = settings.ASSETS_DIR / "svnit_logo.png"
        if svnit_logo.exists():
            p_logo = header_table.rows[0].cells[0].paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.add_run().add_picture(str(svnit_logo), width=Inches(0.9))

        # Center Text
        c_mid = header_table.rows[0].cells[1]
        p1 = c_mid.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run("सरदार वल्लभभाई राष्ट्रीय प्रौद्योगिकी संस्थान, सूरत\n")
        r1.bold = True
        r1.font.size = Pt(9.5)
        
        r2 = p1.add_run("SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT\n")
        r2.bold = True
        r2.font.size = Pt(9.5)

        r3 = p1.add_run("शिक्षा मंत्रालय, भारत सरकार के द्वारा NITSER अधिनियम के तहत स्थापित एक राष्ट्रीय महत्व का संस्थान\n")
        r3.font.size = Pt(7)
        r3.font.color.rgb = RGBColor(71, 85, 105)

        r4 = p1.add_run("An Institute of National Importance, Established under NITSER Act by Ministry of Education, Govt. of India\n")
        r4.font.size = Pt(7)
        r4.font.color.rgb = RGBColor(71, 85, 105)

        r5 = p1.add_run("DEAN STUDENT WELFARE")
        r5.bold = True
        r5.font.size = Pt(9)

        # Right box: FORM BC-R
        c_right = header_table.rows[0].cells[2]
        p_bc = c_right.paragraphs[0]
        p_bc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_bc = p_bc.add_run("FORM\nBC-R")
        r_bc.bold = True
        r_bc.font.size = Pt(10)
        set_cell_background(c_right, "F1F5F9")
        set_cell_margins(c_right, 60, 60, 60, 60)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(4)
        r_t = p_title.add_run("BILL CERTIFICATE\n(RECURRING EXPENDITURE)")
        r_t.bold = True
        r_t.font.size = Pt(11)

        # Metadata table
        meta_rows = [
            (f"Chapter Name: {data.chapter_name}", f"Code No.: {data.code_no}"),
            (f"Approval Note No.: {data.approval_note_no}", f"Date: {data.date}"),
            (f"Purchase Order No. & date: {data.purchase_order_no}", f"Head of Account: {data.head_of_account}"),
            (f"Advance drawn in Rs.: {data.advance_drawn}", ""),
        ]
        if data.has_original_bill_details or any((data.original_approval_sent_with_bill_no, data.original_bill_date, data.original_bill_rs, data.original_bill_ms)):
            meta_rows.extend([
                (f"Original approval sent with bill no.: {data.original_approval_sent_with_bill_no or '-'}", f"Original bill date: {data.original_bill_date or '-'}"),
                (f"Original bill amount (Rs.): {data.original_bill_rs or '-'}", f"Original bill M/s: {data.original_bill_ms or '-'}"),
            ])
        meta_table = doc.add_table(rows=len(meta_rows), cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        mw = [Inches(4.5), Inches(2.7)]
        for r in meta_table.rows:
            r.cells[0].width = mw[0]
            r.cells[1].width = mw[1]

        def set_meta(row, left_txt, right_txt):
            p_l = row.cells[0].paragraphs[0]
            p_l.paragraph_format.space_after = Pt(1)
            p_l.add_run(left_txt).font.size = Pt(8.5)
            p_r = row.cells[1].paragraphs[0]
            p_r.paragraph_format.space_after = Pt(1)
            p_r.add_run(right_txt).font.size = Pt(8.5)

        for idx, (left_text, right_text) in enumerate(meta_rows):
            set_meta(meta_table.rows[idx], left_text, right_text)
        add_table_borders(meta_table)

        # Certification quote
        p_cert = doc.add_paragraph()
        p_cert.paragraph_format.space_before = Pt(6)
        p_cert.paragraph_format.space_after = Pt(6)
        r_c = p_cert.add_run(
            f'“I, am {data.certifier_name}, personally satisfied that the goods (described below) purchased are '
            f'of the requisite quality and specification and have been purchased from a reliable supplier / '
            f'contractor at a reasonable price.”'
        )
        r_c.font.size = Pt(8.5)
        r_c.italic = True

        # Items Table
        items_table = doc.add_table(rows=len(calc_items) + 4, cols=5)
        items_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        items_table.autofit = False
        col_widths = [Inches(0.6), Inches(3.6), Inches(0.8), Inches(1.1), Inches(1.1)]
        for row in items_table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

        # Table Header
        headers = ["Sr. No.", "Name of the Equipment/ item(s)", "Quantity", "Unit Cost (Rs.)", "Total Cost (Rs.)"]
        for i, h in enumerate(headers):
            c = items_table.rows[0].cells[i]
            set_cell_background(c, "E2E8F0")
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(8.5)

        # Table Items
        for idx, itm in enumerate(calc_items, start=1):
            row = items_table.rows[idx]
            vals = [str(itm["sr_no"]), itm["name"], str(itm["quantity"]), f"{itm['unit_cost']:.2f}", f"{itm['total_cost']:.2f}"]
            for c_i, v in enumerate(vals):
                cell = row.cells[c_i]
                set_cell_margins(cell, 60, 60, 70, 70)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if c_i in [0, 2]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif c_i in [3, 4]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.add_run(v).font.size = Pt(8.5)

        # Total Rows
        row_sub = items_table.rows[len(calc_items) + 1]
        row_sub.cells[0].merge(row_sub.cells[3])
        p_sub = row_sub.cells[0].paragraphs[0]
        p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sub.add_run("Total Amount (Rs.):").bold = True
        p_sub.runs[0].font.size = Pt(8.5)
        p_sub_val = row_sub.cells[4].paragraphs[0]
        p_sub_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sub_val.add_run(f"{subtotal:.2f}").bold = True
        p_sub_val.runs[0].font.size = Pt(8.5)

        row_tax = items_table.rows[len(calc_items) + 2]
        row_tax.cells[0].merge(row_tax.cells[3])
        p_tax = row_tax.cells[0].paragraphs[0]
        p_tax.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        tax_lbl = f"+ Taxes ({data.tax_details}):" if getattr(data, "tax_details", None) else "+ Taxes or other charges:"
        p_tax.add_run(tax_lbl).bold = True
        p_tax.runs[0].font.size = Pt(8.5)
        p_tax_val = row_tax.cells[4].paragraphs[0]
        p_tax_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_tax_val.add_run(f"{taxes:.2f}").font.size = Pt(8.5)

        row_grand = items_table.rows[len(calc_items) + 3]
        row_grand.cells[0].merge(row_grand.cells[3])
        set_cell_background(row_grand.cells[0], "F1F5F9")
        set_cell_background(row_grand.cells[4], "F1F5F9")
        p_g = row_grand.cells[0].paragraphs[0]
        p_g.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_g.add_run("Total Estimated Amount Including Taxes (Rs.):").bold = True
        p_g.runs[0].font.size = Pt(8.5)
        p_g_val = row_grand.cells[4].paragraphs[0]
        p_g_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_g_val.add_run(f"{grand_total:.2f}").bold = True
        p_g_val.runs[0].font.size = Pt(8.5)

        add_table_borders(items_table)

        # Supplier & Payment Details side-by-side or stacked
        if data.has_advance_adjustment or data.advance_adjustment_notes:
            adjustment = doc.add_table(rows=1, cols=1)
            adjustment.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_cell_margins(adjustment.rows[0].cells[0], 70, 70, 90, 90)
            adjustment.rows[0].cells[0].paragraphs[0].add_run(
                f"Advance Adjustment: {data.advance_adjustment_notes or '-'}"
            ).font.size = Pt(8)
            add_table_borders(adjustment)

        doc.add_paragraph().paragraph_format.space_before = Pt(4)
        
        details_table = doc.add_table(rows=1, cols=2)
        details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        details_table.autofit = False
        details_table.rows[0].cells[0].width = Inches(3.6)
        details_table.rows[0].cells[1].width = Inches(3.6)

        c_sup = details_table.rows[0].cells[0]
        set_cell_background(c_sup, "FAFAFA")
        set_cell_margins(c_sup, 80, 80, 100, 100)
        p_sup = c_sup.paragraphs[0]
        p_sup.add_run("SUPPLIER DETAILS\n").bold = True
        p_sup.runs[0].font.size = Pt(9)
        sup_text = (
            f"Name of the Supplier: {data.supplier_name}\n"
            f"GST No.: {data.supplier_gst or '-'}\n"
            f"Any other Taxes: {data.any_other_taxes or '-'}\n"
            f"Invoice No. & Date: {data.invoice_no_and_date or '-'}\n"
            f"Amount: Rs. {grand_total:.2f}"
        )
        p_sup.add_run(sup_text).font.size = Pt(8)

        c_pay = details_table.rows[0].cells[1]
        set_cell_background(c_pay, "FAFAFA")
        set_cell_margins(c_pay, 80, 80, 100, 100)
        p_pay = c_pay.paragraphs[0]
        p_pay.add_run("PAYMENT DETAILS\n").bold = True
        p_pay.runs[0].font.size = Pt(9)
        pay_text = (
            f"Name of the Party: {data.party_name or data.supplier_name}\n"
            f"Account Number: {data.account_number or '-'}\n"
            f"Account Holder: {data.account_holder_name or '-'}\n"
            f"Bank Name: {data.bank_name or '-'}\n"
            f"Branch Name: {data.branch_name or '-'}\n"
            f"IFSC Code: {data.ifsc_code or '-'}"
        )
        p_pay.add_run(pay_text).font.size = Pt(8)

        add_table_borders(details_table)

        # PAGE 2: Stock Entry & Signatures
        doc.add_page_break()

        p_pg2_title = doc.add_paragraph()
        p_pg2_title.add_run("STOCK ENTRY DETAILS").bold = True
        p_pg2_title.runs[0].font.size = Pt(10)
        
        p_lab = doc.add_paragraph()
        p_lab.add_run(f"Name of the Laboratory / Office: {data.lab_office}").font.size = Pt(9)

        # Stock table
        stk_table = doc.add_table(rows=5, cols=4)
        stk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        stk_table.autofit = False
        stk_widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8)]
        for r in stk_table.rows:
            for ci, w in enumerate(stk_widths):
                r.cells[ci].width = w

        stk_headers = ["", "Recurring Register", "Consumable Register", "Miscellaneous Register"]
        for ci, h in enumerate(stk_headers):
            c = stk_table.rows[0].cells[ci]
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(h).bold = True
            p.runs[0].font.size = Pt(8)

        chosen_register = data.register_type or "Recurring Register"
        stk_rows = [
            ("Register Type", "[X]" if chosen_register == "Recurring Register" else "[ ]", "[X]" if chosen_register == "Consumable Register" else "[ ]", "[X]" if chosen_register == "Miscellaneous Register" else "[ ]"),
            ("Serial Number", data.serial_number or "-", "-", "-"),
            ("Page Number", data.page_number or "-", "-", "-"),
            ("Date & Signature", f"Date: {data.stock_date or data.date}\nClerk: {data.clerk_name or ''}" if chosen_register == "Recurring Register" else "", f"Date: {data.stock_date or data.date}\nClerk: {data.clerk_name or ''}" if chosen_register == "Consumable Register" else "", f"Date: {data.stock_date or data.date}\nClerk: {data.clerk_name or ''}" if chosen_register == "Miscellaneous Register" else "")
        ]
        for ri, row_data in enumerate(stk_rows, start=1):
            for ci, txt in enumerate(row_data):
                c = stk_table.rows[ri].cells[ci]
                p = c.paragraphs[0]
                p.add_run(txt).font.size = Pt(8)

        add_table_borders(stk_table)

        # Indenter Details
        p_ind = doc.add_paragraph()
        p_ind.paragraph_format.space_before = Pt(12)
        p_ind.add_run("INDENTER DETAILS").bold = True
        p_ind.runs[0].font.size = Pt(10)

        ind_table = doc.add_table(rows=1, cols=2)
        ind_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        ind_table.autofit = False
        ind_table.rows[0].cells[0].width = Inches(4.5)
        ind_table.rows[0].cells[1].width = Inches(2.7)

        c_ind_l = ind_table.rows[0].cells[0]
        p_ind_l = c_ind_l.paragraphs[0]
        p_ind_l.add_run(
            f"Name: {data.indenter_name}\n"
            f"Designation: {data.indenter_designation}\n"
            f"Department: {data.indenter_department}"
        ).font.size = Pt(8.5)

        c_ind_r = ind_table.rows[0].cells[1]
        p_ind_r = c_ind_r.paragraphs[0]
        p_ind_r.paragraph_format.space_before = Pt(16)
        p_ind_r.add_run("Signature with Date\n___________________").font.size = Pt(8.5)

        # Recommendation (HoD)
        p_rec = doc.add_paragraph()
        p_rec.paragraph_format.space_before = Pt(10)
        p_rec.add_run("Recommendation (Head of the Department)\n").bold = True
        p_rec.runs[0].font.size = Pt(9.5)
        p_rec.add_run(
            f"Recommended Amount: Rs. {float(data.recommended_amount) if data.recommended_amount is not None else grand_total:.2f}\n"
            "Approved and Recommended for Bill Payment\n\n"
            "Date: _______________                       Dean Student Welfare: _______________"
        ).font.size = Pt(8.5)

        # Accounts Section
        p_acc = doc.add_paragraph()
        p_acc.paragraph_format.space_before = Pt(10)
        p_acc.add_run("(FOR USE IN ACCOUNT SECTION)").bold = True
        p_acc.runs[0].font.size = Pt(9.5)

        acc_box = doc.add_table(rows=1, cols=1)
        acc_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        acc_box.rows[0].cells[0].width = Inches(7.2)
        c_acc = acc_box.rows[0].cells[0]
        set_cell_margins(c_acc, 80, 80, 80, 80)
        p_ab = c_acc.paragraphs[0]
        payment_options = ["Cash", "Cheque", "Online", "Adjusted Against Advance"]
        payment_modes = "  ".join(
            f"{'[X]' if mode in data.payment_modes else '[ ]'} {mode}" for mode in payment_options
        )
        p_ab.add_run(
            f"Date: _______________   Head of Account: {data.head_of_account}   Budget Code No.: {data.code_no}\n"
            f"Mode of Payment: {payment_modes}\n"
            f"Details of Mode of Payment: {data.payment_mode_details or '__________________________________________'}\n"
            f"Passed for payment/adjustment of Amount in (Rs.): {grand_total:.2f}\n\n"
            "Signatures:\n\n"
            "Office Supt.                    Dy. Registrar (A/cs.)                    Registrar/Director"
        ).font.size = Pt(8)
        add_table_borders(acc_box)

        doc.save(str(output_path))
        return output_path

    @classmethod
    def generate_pdf(cls, data: BillCertificateRequest, output_path: Path):
        calc_items, subtotal, taxes, grand_total = cls.calculate_totals(data)
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=30,
            rightMargin=30,
            topMargin=26,
            bottomMargin=26
        )
        
        styles = getSampleStyleSheet()
        normal = styles["Normal"]

        inst_eng = ParagraphStyle("InstEng", parent=normal, fontName="Helvetica-Bold", fontSize=8.5, leading=10, alignment=1)
        inst_sub = ParagraphStyle("InstSub", parent=normal, fontName="Helvetica", fontSize=6.5, leading=8.5, alignment=1, textColor=colors.HexColor("#475569"))
        bcr_style = ParagraphStyle("BCR", parent=normal, fontName="Helvetica-Bold", fontSize=9, leading=11, alignment=1)
        
        title_style = ParagraphStyle("Title", parent=normal, fontName="Helvetica-Bold", fontSize=11, leading=13, alignment=1)
        body = ParagraphStyle("Body", parent=normal, fontName="Helvetica", fontSize=8, leading=10.5)
        body_bold = ParagraphStyle("BodyBold", parent=normal, fontName="Helvetica-Bold", fontSize=8, leading=10.5)
        body_right = ParagraphStyle("BodyRight", parent=body, alignment=2)
        body_right_bold = ParagraphStyle("BodyRightBold", parent=body_bold, alignment=2)
        body_center = ParagraphStyle("BodyCenter", parent=body, alignment=1)

        story = []

        # PAGE 1 Header
        svnit_logo = settings.ASSETS_DIR / "svnit_logo.png"
        logo_left = RLImage(str(svnit_logo), width=0.75*inch, height=0.72*inch) if svnit_logo.exists() else ""
        
        inst_paragraphs = [
            Paragraph("SARDAR VALLABHBHAI NATIONAL INSTITUTE OF TECHNOLOGY, SURAT", inst_eng),
            Spacer(1, 2),
            Paragraph("An Institute of National Importance, Established under NITSER Act by Ministry of Education, Govt. of India", inst_sub),
            Spacer(1, 2),
            Paragraph("<b>DEAN STUDENT WELFARE</b>", inst_eng),
        ]

        bcr_box = [Paragraph("<b>FORM<br/>BC-R</b>", bcr_style)]

        head_table = Table([[logo_left, inst_paragraphs, bcr_box]], colWidths=[0.85*inch, 5.7*inch, 0.85*inch])
        head_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('BOX', (2,0), (2,0), 1, colors.HexColor("#64748b")),
            ('BACKGROUND', (2,0), (2,0), colors.HexColor("#f8fafc")),
            ('TOPPADDING', (2,0), (2,0), 6),
            ('BOTTOMPADDING', (2,0), (2,0), 6),
        ]))
        story.append(head_table)
        story.append(Spacer(1, 4))

        # Title
        story.append(Paragraph("BILL CERTIFICATE<br/>(RECURRING EXPENDITURE)", title_style))
        story.append(Spacer(1, 6))

        # Metadata
        meta_data = [
            [Paragraph(f"<b>Chapter Name:</b> {escape(data.chapter_name)}", body), Paragraph(f"<b>Code No.:</b> {escape(data.code_no)}", body)],
            [Paragraph(f"<b>Approval Note No.:</b> {escape(data.approval_note_no)}", body), Paragraph(f"<b>Date:</b> {escape(data.date)}", body)],
            [Paragraph(f"<b>Purchase Order No. &amp; date:</b> {escape(data.purchase_order_no)}", body), Paragraph(f"<b>Head of Account:</b> {escape(data.head_of_account)}", body)],
            [Paragraph(f"<b>Advance drawn in Rs.:</b> {escape(data.advance_drawn)}", body), Paragraph("", body)]
        ]
        if data.has_original_bill_details or any((data.original_approval_sent_with_bill_no, data.original_bill_date, data.original_bill_rs, data.original_bill_ms)):
            meta_data.extend([
                [Paragraph(f"<b>Original approval sent with bill no.:</b> {escape(data.original_approval_sent_with_bill_no or '-')}", body), Paragraph(f"<b>Original bill date:</b> {escape(data.original_bill_date or '-')}", body)],
                [Paragraph(f"<b>Original bill amount (Rs.):</b> {escape(data.original_bill_rs or '-')}", body), Paragraph(f"<b>Original bill M/s:</b> {escape(data.original_bill_ms or '-')}", body)],
            ])
        meta_tbl = Table(meta_data, colWidths=[4.6*inch, 2.8*inch])
        meta_tbl.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 1),
            ('BOTTOMPADDING', (0,0), (-1,-1), 1),
            ('LEFTPADDING', (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(meta_tbl)
        story.append(Spacer(1, 4))

        # Certificate statement
        cert_text = (
            f'<i>“I, am {escape(data.certifier_name)}, personally satisfied that the goods (described below) purchased are '
            f'of the requisite quality and specification and have been purchased from a reliable supplier / '
            f'contractor at a reasonable price.”</i>'
        )
        story.append(Table([[Paragraph(cert_text, body)]], colWidths=[7.4*inch], style=TableStyle([
            ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#94a3b8")),
            ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ])))
        story.append(Spacer(1, 5))

        # Items Table
        item_rows = [
            [
                Paragraph("<b>Sr. No.</b>", body_center),
                Paragraph("<b>Name of the Equipment/ item(s)</b>", body_bold),
                Paragraph("<b>Quantity</b>", body_center),
                Paragraph("<b>Unit Cost (Rs.)</b>", body_right_bold),
                Paragraph("<b>Total Cost (Rs.)</b>", body_right_bold)
            ]
        ]
        for itm in calc_items:
            item_rows.append([
                Paragraph(str(itm["sr_no"]), body_center),
                Paragraph(escape(itm["name"]), body),
                Paragraph(str(itm["quantity"]), body_center),
                Paragraph(f"{itm['unit_cost']:.2f}", body_right),
                Paragraph(f"{itm['total_cost']:.2f}", body_right),
            ])

        # Summary rows
        item_rows.append([
            "", Paragraph("<b>Total Amount (Rs.)</b>", body_right_bold), "", "",
            Paragraph(f"<b>{subtotal:.2f}</b>", body_right_bold)
        ])
        pdf_tax_lbl = f"<b>+ Taxes ({data.tax_details})</b>" if getattr(data, "tax_details", None) else "<b>+ Taxes or other charges</b>"
        item_rows.append([
            "", Paragraph(pdf_tax_lbl, body_right_bold), "", "",
            Paragraph(f"{taxes:.2f}", body_right)
        ])
        item_rows.append([
            "", Paragraph("<b>Total Estimated Amount Including Taxes (Rs.)</b>", body_right_bold), "", "",
            Paragraph(f"<b>{grand_total:.2f}</b>", body_right_bold)
        ])

        total_row_start = len(calc_items) + 1
        items_pdf_table = Table(item_rows, colWidths=[0.55*inch, 3.85*inch, 0.75*inch, 1.1*inch, 1.15*inch])
        items_pdf_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#e2e8f0")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('SPAN', (1, total_row_start), (3, total_row_start)),
            ('SPAN', (1, total_row_start + 1), (3, total_row_start + 1)),
            ('SPAN', (1, total_row_start + 2), (3, total_row_start + 2)),
            ('BACKGROUND', (0, total_row_start + 2), (-1, total_row_start + 2), colors.HexColor("#f1f5f9")),
        ]))
        story.append(items_pdf_table)
        story.append(Spacer(1, 6))

        # Supplier & Payment Side-by-side
        sup_lines = [
            "<b>SUPPLIER DETAILS</b>",
            f"Name: {escape(data.supplier_name)}",
            f"GST No.: {escape(data.supplier_gst or '-')} ",
            f"Any other Taxes: {escape(data.any_other_taxes or '-')}",
            f"Invoice No. & Date: {escape(data.invoice_no_and_date or '-')}",
            f"Amount: Rs. {grand_total:.2f}"
        ]
        pay_lines = [
            "<b>PAYMENT DETAILS</b>",
            f"Party Name: {escape(data.party_name or data.supplier_name)}",
            f"Account Number: {escape(data.account_number or '-')}",
            f"Account Holder: {escape(data.account_holder_name or '-')}",
            f"Bank: {escape(data.bank_name or '-')}",
            f"Branch: {escape(data.branch_name or '-')}",
            f"IFSC Code: {escape(data.ifsc_code or '-')}"
        ]

        details_pdf = Table([
            [Paragraph("<br/>".join(sup_lines), body), Paragraph("<br/>".join(pay_lines), body)]
        ], colWidths=[3.7*inch, 3.7*inch])
        details_pdf.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#fafafa")),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING', (0,0), (-1,-1), 6),
            ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(details_pdf)
        if data.has_advance_adjustment or data.advance_adjustment_notes:
            story.append(Spacer(1, 4))
            adjustment_pdf = Table([[Paragraph(f"<b>Advance Adjustment:</b> {escape(data.advance_adjustment_notes or '-')}", body)]], colWidths=[7.4*inch])
            adjustment_pdf.setStyle(TableStyle([
                ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#94a3b8")),
                ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6),
                ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ]))
            story.append(adjustment_pdf)

        # PAGE 2
        story.append(PageBreak())

        story.append(Paragraph("<b>STOCK ENTRY DETAILS</b>", body_bold))
        story.append(Spacer(1, 2))
        story.append(Paragraph(f"Name of the Laboratory / Office: <b>{escape(data.lab_office)}</b>", body))
        story.append(Spacer(1, 4))

        chosen_register = data.register_type or "Recurring Register"
        register_markers = ["[X]" if chosen_register == register_name else "[ ]" for register_name in ("Recurring Register", "Consumable Register", "Miscellaneous Register")]
        stock_details = f"Date: {escape(data.stock_date or data.date)}<br/>Staff: {escape(data.clerk_name or '')}"
        stock_data_cells = [stock_details if chosen_register == register_name else "" for register_name in ("Recurring Register", "Consumable Register", "Miscellaneous Register")]
        stk_data = [
            [
                Paragraph("", body),
                Paragraph("<b>Recurring Register</b>", body_center),
                Paragraph("<b>Consumable Register</b>", body_center),
                Paragraph("<b>Miscellaneous Register</b>", body_center)
            ],
            [Paragraph("Selected Register", body_bold), *[Paragraph(marker, body_center) for marker in register_markers]],
            [Paragraph("Serial Number", body_bold), *[Paragraph(escape(data.serial_number or "-"), body_center) if chosen_register == register_name else Paragraph("-", body_center) for register_name in ("Recurring Register", "Consumable Register", "Miscellaneous Register")]],
            [Paragraph("Page Number", body_bold), *[Paragraph(escape(data.page_number or "-"), body_center) if chosen_register == register_name else Paragraph("-", body_center) for register_name in ("Recurring Register", "Consumable Register", "Miscellaneous Register")]],
            [Paragraph("Date & Signature", body_bold), *[Paragraph(cell, body) for cell in stock_data_cells]],
        ]
        stk_table_pdf = Table(stk_data, colWidths=[1.8*inch, 1.85*inch, 1.85*inch, 1.9*inch])
        stk_table_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,1), colors.HexColor("#f1f5f9")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(stk_table_pdf)
        story.append(Spacer(1, 10))

        # Indenter Details
        story.append(Paragraph("<b>INDENTER DETAILS</b>", body_bold))
        story.append(Spacer(1, 2))
        ind_data = [
            [
                Paragraph(f"Name: <b>{escape(data.indenter_name)}</b><br/>Designation: {escape(data.indenter_designation)}<br/>Department: {escape(data.indenter_department)}", body),
                Paragraph("<br/><br/>Signature with Date: ___________________", body)
            ]
        ]
        ind_tbl = Table(ind_data, colWidths=[4.6*inch, 2.8*inch])
        ind_tbl.setStyle(TableStyle([('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#94a3b8")), ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")), ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6), ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5)]))
        story.append(ind_tbl)
        story.append(Spacer(1, 10))

        # Recommendation (HoD)
        story.append(Paragraph("<b>Recommendation (Head of the Department)</b>", body_bold))
        story.append(Spacer(1, 2))
        rec_data = [
            [
                Paragraph(f"Recommended Amount: <b>Rs. {float(data.recommended_amount) if data.recommended_amount is not None else grand_total:.2f}</b><br/>Approved and Recommended for Bill Payment<br/><br/>Date: _______________", body),
                Paragraph("<br/><br/>Dean Student Welfare: ___________________", body)
            ]
        ]
        rec_tbl = Table(rec_data, colWidths=[4.6*inch, 2.8*inch])
        rec_tbl.setStyle(TableStyle([('BOX', (0,0), (-1,-1), 0.8, colors.HexColor("#94a3b8")), ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")), ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6), ('TOPPADDING', (0,0), (-1,-1), 6), ('BOTTOMPADDING', (0,0), (-1,-1), 6)]))
        story.append(rec_tbl)
        story.append(Spacer(1, 10))

        # Account section
        story.append(Paragraph("<b>(FOR USE IN ACCOUNT SECTION)</b>", body_bold))
        story.append(Spacer(1, 3))
        mode_options = ["Cash", "Cheque", "Online", "Adjusted Against Advance"]
        mode_choices = "   ".join(f"{'[X]' if mode in data.payment_modes else '[ ]'} {mode}" for mode in mode_options)
        acc_data = [
            [Paragraph("<b>Date</b>", body_bold), Paragraph("________________", body), Paragraph("<b>Head of Account</b>", body_bold), Paragraph(escape(data.head_of_account), body)],
            [Paragraph("<b>Budget Code No.</b>", body_bold), Paragraph(escape(data.code_no), body), Paragraph("<b>Amount Passed (Rs.)</b>", body_bold), Paragraph(f"{grand_total:.2f}", body_right_bold)],
            [Paragraph("<b>Mode of Payment</b>", body_bold), Paragraph(escape(mode_choices), body), "", ""],
            [Paragraph("<b>Payment Details</b>", body_bold), Paragraph(escape(data.payment_mode_details or "_______________________________________________"), body), "", ""],
            [Paragraph("<b>Signatures</b>", body_bold), Paragraph("Office Supt.", body_center), Paragraph("Dy. Registrar (A/cs.)", body_center), Paragraph("Registrar / Director", body_center)],
            [Paragraph("", body), Paragraph("________________", body_center), Paragraph("________________", body_center), Paragraph("________________", body_center)],
        ]
        acc_box_pdf = Table(acc_data, colWidths=[1.35*inch, 2.05*inch, 1.65*inch, 2.35*inch])
        acc_box_pdf.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('SPAN', (1,2), (3,2)),
            ('SPAN', (1,3), (3,3)),
            ('SPAN', (1,4), (1,4)),
            ('SPAN', (2,4), (2,4)),
            ('SPAN', (3,4), (3,4)),
            ('SPAN', (0,5), (0,5)),
            ('SPAN', (1,5), (1,5)),
            ('SPAN', (2,5), (2,5)),
            ('SPAN', (3,5), (3,5)),
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#fcfcfc")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(acc_box_pdf)

        doc.build(story)
        return output_path
