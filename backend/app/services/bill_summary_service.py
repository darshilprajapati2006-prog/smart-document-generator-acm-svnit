from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)

from app.schemas.bill_summary import BillSummaryRequest

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_table_borders(table, border_color="94A3B8"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

class BillSummaryService:
    @staticmethod
    def calculate_totals(data: BillSummaryRequest):
        calc_rows = []
        total_prize = 0.0
        total_other = 0.0
        for idx, row in enumerate(data.rows, start=1):
            pm = round(float(row.prize_money), 2)
            oth = round(float(row.other_amount), 2)
            total_prize += pm
            total_other += oth
            calc_rows.append({
                "sr_no": idx,
                "name": row.name,
                "account_details": row.account_details or (f"A/C: {row.account_no}" if row.account_no else "-"),
                "to_be_paid_to": row.to_be_paid_to or row.name,
                "prize_money": pm,
                "other_amount": oth,
                "total": round(pm + oth, 2),
                "account_no": row.account_no or "-",
                "ifsc_code": row.ifsc_code or "-",
            })
        total_prize = round(total_prize, 2)
        total_other = round(total_other, 2)
        grand_total = round(total_prize + total_other, 2)
        return calc_rows, total_prize, total_other, grand_total

    @classmethod
    def generate_docx(cls, data: BillSummaryRequest, output_path: Path):
        calc_rows, total_prize, total_other, grand_total = cls.calculate_totals(data)
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # Header Title
        p_h1 = doc.add_paragraph()
        p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h1.paragraph_format.space_after = Pt(2)
        r_h1 = p_h1.add_run("Association for Computing Machinery")
        r_h1.bold = True
        r_h1.font.size = Pt(13)

        p_h2 = doc.add_paragraph()
        p_h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h2.paragraph_format.space_after = Pt(8)
        r_h2 = p_h2.add_run("Sardar Vallabhbhai National Institute of Technology, Surat - 395007")
        r_h2.font.size = Pt(9.5)
        r_h2.font.color.rgb = RGBColor(71, 85, 105)

        # Date right aligned
        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_date.paragraph_format.space_after = Pt(10)
        r_date = p_date.add_run(f"Date: {data.date}")
        r_date.bold = True
        r_date.font.size = Pt(9.5)

        # Document Title & Subtitle
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(2)
        r_t = p_title.add_run(data.title.upper())
        r_t.bold = True
        r_t.font.size = Pt(12)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(12)
        r_s = p_sub.add_run(data.subtitle)
        r_s.italic = True
        r_s.font.size = Pt(9.5)

        # Table 1: Summary of Reimbursements
        t1 = doc.add_table(rows=len(calc_rows) + 2, cols=6)
        t1.alignment = WD_TABLE_ALIGNMENT.CENTER
        t1.autofit = False
        widths = [Inches(0.6), Inches(1.8), Inches(1.8), Inches(1.4), Inches(1.0), Inches(0.8)]
        for r in t1.rows:
            for ci, w in enumerate(widths):
                r.cells[ci].width = w

        headers = ["Sr. No.", "Name", "Account Details", "To Be Paid to", "Prize Money (INR)", "Other"]
        for ci, h in enumerate(headers):
            c = t1.rows[0].cells[ci]
            set_cell_background(c, "E2E8F0")
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(8.5)

        for ri, rdata in enumerate(calc_rows, start=1):
            row = t1.rows[ri]
            vals = [
                str(rdata["sr_no"]),
                rdata["name"],
                rdata["account_details"],
                rdata["to_be_paid_to"],
                f"{rdata['prize_money']:.2f}" if rdata['prize_money'] else "-",
                f"{rdata['other_amount']:.2f}" if rdata['other_amount'] else "-"
            ]
            for ci, val in enumerate(vals):
                c = row.cells[ci]
                set_cell_margins(c, 60, 60, 70, 70)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if ci in [0]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif ci in [4, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.add_run(val).font.size = Pt(8.5)

        # Grand Total Row
        row_gt = t1.rows[len(calc_rows) + 1]
        row_gt.cells[0].merge(row_gt.cells[3])
        set_cell_background(row_gt.cells[0], "F1F5F9")
        set_cell_background(row_gt.cells[4], "F1F5F9")
        set_cell_background(row_gt.cells[5], "F1F5F9")
        
        p_gt_lbl = row_gt.cells[0].paragraphs[0]
        p_gt_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_gt_lbl.add_run("Grand Total:").bold = True
        p_gt_lbl.runs[0].font.size = Pt(9)

        p_gt_p = row_gt.cells[4].paragraphs[0]
        p_gt_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_gt_p.add_run(f"Rs. {total_prize:.2f}").bold = True
        p_gt_p.runs[0].font.size = Pt(8.5)

        p_gt_o = row_gt.cells[5].paragraphs[0]
        p_gt_o.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_gt_o.add_run(f"Rs. {total_other:.2f}").bold = True
        p_gt_o.runs[0].font.size = Pt(8.5)

        add_table_borders(t1)

        # Signatures 3 Columns
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(36)

        sig_table = doc.add_table(rows=1, cols=3)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = False
        sw = [Inches(2.4), Inches(2.4), Inches(2.4)]
        for ci, w in enumerate(sw):
            sig_table.rows[0].cells[ci].width = w

        def set_sig_col(cell, role, name, club):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            r_role = p.add_run(f"{role}\n")
            r_role.bold = True
            r_role.font.size = Pt(9)
            p.add_run(f"{name}\n").font.size = Pt(9)
            p.add_run(f"{club}\n\n").font.size = Pt(8.5)
            r_line = p.add_run("Signature: ______________")
            r_line.font.size = Pt(8)
            r_line.font.color.rgb = RGBColor(148, 163, 184)

        set_sig_col(sig_table.rows[0].cells[0], data.student_head_title, data.student_head_name, data.club_name)
        set_sig_col(sig_table.rows[0].cells[1], data.student_secretary_title, data.student_secretary_name, data.club_name)
        set_sig_col(sig_table.rows[0].cells[2], data.faculty_chairman_title, data.faculty_chairman_name, data.club_name)

        # PAGE 2: Other Bank Account Details
        doc.add_page_break()

        p_pg2_title = doc.add_paragraph()
        p_pg2_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pg2_title.paragraph_format.space_before = Pt(10)
        p_pg2_title.paragraph_format.space_after = Pt(14)
        r2 = p_pg2_title.add_run("Other Bank Account Details")
        r2.bold = True
        r2.font.size = Pt(12)

        # Bank table
        t2 = doc.add_table(rows=len(calc_rows) + 2, cols=5)
        t2.alignment = WD_TABLE_ALIGNMENT.CENTER
        t2.autofit = False
        t2_widths = [Inches(0.6), Inches(2.3), Inches(2.3), Inches(1.1), Inches(1.1)]
        for r in t2.rows:
            for ci, w in enumerate(t2_widths):
                r.cells[ci].width = w

        t2_headers = ["Sr. No.", "Name", "Account No.", "IFSC Code", "Amount"]
        for ci, h in enumerate(t2_headers):
            c = t2.rows[0].cells[ci]
            set_cell_background(c, "E2E8F0")
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(8.5)

        for ri, rdata in enumerate(calc_rows, start=1):
            row = t2.rows[ri]
            vals = [
                str(rdata["sr_no"]),
                rdata["name"],
                rdata["account_no"],
                rdata["ifsc_code"],
                f"{rdata['total']:.2f}"
            ]
            for ci, val in enumerate(vals):
                c = row.cells[ci]
                set_cell_margins(c, 60, 60, 70, 70)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if ci == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif ci == 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.add_run(val).font.size = Pt(8.5)

        # Total row
        row_t2_total = t2.rows[len(calc_rows) + 1]
        row_t2_total.cells[0].merge(row_t2_total.cells[3])
        set_cell_background(row_t2_total.cells[0], "F1F5F9")
        set_cell_background(row_t2_total.cells[4], "F1F5F9")
        
        p_t2_lbl = row_t2_total.cells[0].paragraphs[0]
        p_t2_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_t2_lbl.add_run("Total:").bold = True
        p_t2_lbl.runs[0].font.size = Pt(9)

        p_t2_val = row_t2_total.cells[4].paragraphs[0]
        p_t2_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_t2_val.add_run(f"Rs. {grand_total:.2f}").bold = True
        p_t2_val.runs[0].font.size = Pt(9)

        add_table_borders(t2)

        # Page 2 signature: Faculty Chairman
        p_sig2 = doc.add_paragraph()
        p_sig2.paragraph_format.space_before = Pt(40)
        r_f = p_sig2.add_run(
            f"{data.faculty_chairman_title}\n"
            f"{data.faculty_chairman_name}\n"
            f"{data.club_name}\n\n"
            "Signature: ____________________"
        )
        r_f.font.size = Pt(9)

        doc.save(str(output_path))
        return output_path

    @classmethod
    def generate_pdf(cls, data: BillSummaryRequest, output_path: Path):
        calc_rows, total_prize, total_other, grand_total = cls.calculate_totals(data)
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        normal = styles["Normal"]

        inst_style = ParagraphStyle("Inst", parent=normal, fontName="Helvetica-Bold", fontSize=12, leading=15, alignment=1)
        subinst_style = ParagraphStyle("SubInst", parent=normal, fontName="Helvetica", fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#475569"))
        date_style = ParagraphStyle("Date", parent=normal, fontName="Helvetica-Bold", fontSize=9, leading=12, alignment=2)
        title_style = ParagraphStyle("Title", parent=normal, fontName="Helvetica-Bold", fontSize=12, leading=15, alignment=1)
        subtitle_style = ParagraphStyle("Sub", parent=normal, fontName="Helvetica-Oblique", fontSize=9.5, leading=13, alignment=1)

        body = ParagraphStyle("Body", parent=normal, fontName="Helvetica", fontSize=8, leading=10.5)
        body_bold = ParagraphStyle("BodyBold", parent=normal, fontName="Helvetica-Bold", fontSize=8, leading=10.5)
        body_center = ParagraphStyle("BodyCenter", parent=body, alignment=1)
        body_right = ParagraphStyle("BodyRight", parent=body, alignment=2)
        body_right_bold = ParagraphStyle("BodyRightBold", parent=body_bold, alignment=2)

        story = []

        # PAGE 1
        story.append(Paragraph("Association for Computing Machinery", inst_style))
        story.append(Spacer(1, 2))
        story.append(Paragraph("Sardar Vallabhbhai National Institute of Technology, Surat - 395007", subinst_style))
        story.append(Spacer(1, 6))

        story.append(Paragraph(f"<b>Date:</b> {data.date}", date_style))
        story.append(Spacer(1, 8))

        story.append(Paragraph(f"<b>{data.title.upper()}</b>", title_style))
        story.append(Spacer(1, 3))
        story.append(Paragraph(data.subtitle, subtitle_style))
        story.append(Spacer(1, 10))

        # Main Table
        table_rows = [
            [
                Paragraph("<b>Sr. No.</b>", body_center),
                Paragraph("<b>Name</b>", body_bold),
                Paragraph("<b>Account Details</b>", body_bold),
                Paragraph("<b>To Be Paid to</b>", body_bold),
                Paragraph("<b>Prize Money (INR)</b>", body_right_bold),
                Paragraph("<b>Other</b>", body_right_bold)
            ]
        ]
        for itm in calc_rows:
            table_rows.append([
                Paragraph(str(itm["sr_no"]), body_center),
                Paragraph(itm["name"], body),
                Paragraph(itm["account_details"], body),
                Paragraph(itm["to_be_paid_to"], body),
                Paragraph(f"{itm['prize_money']:.2f}" if itm['prize_money'] else "-", body_right),
                Paragraph(f"{itm['other_amount']:.2f}" if itm['other_amount'] else "-", body_right)
            ])

        # Grand Total row
        table_rows.append([
            "", Paragraph("<b>Grand Total</b>", body_right_bold), "", "",
            Paragraph(f"<b>Rs. {total_prize:.2f}</b>", body_right_bold),
            Paragraph(f"<b>Rs. {total_other:.2f}</b>", body_right_bold)
        ])

        gt_idx = len(calc_rows) + 1
        t1_pdf = Table(table_rows, colWidths=[0.55*inch, 1.8*inch, 1.8*inch, 1.35*inch, 1.0*inch, 0.9*inch])
        t1_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#e2e8f0")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('SPAN', (1, gt_idx), (3, gt_idx)),
            ('BACKGROUND', (0, gt_idx), (-1, gt_idx), colors.HexColor("#f1f5f9")),
        ]))
        story.append(t1_pdf)
        story.append(Spacer(1, 40))

        # Signatures 3 columns
        sig_data = [
            [
                Paragraph(f"<b>{data.student_head_title}</b><br/>{data.student_head_name}<br/>{data.club_name}<br/><br/><font color='#94a3b8'>Signature: _________</font>", body_center),
                Paragraph(f"<b>{data.student_secretary_title}</b><br/>{data.student_secretary_name}<br/>{data.club_name}<br/><br/><font color='#94a3b8'>Signature: _________</font>", body_center),
                Paragraph(f"<b>{data.faculty_chairman_title}</b><br/>{data.faculty_chairman_name}<br/>{data.club_name}<br/><br/><font color='#94a3b8'>Signature: _________</font>", body_center)
            ]
        ]
        sig_tbl_pdf = Table(sig_data, colWidths=[2.45*inch, 2.45*inch, 2.5*inch])
        sig_tbl_pdf.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER')
        ]))
        story.append(KeepTogether(sig_tbl_pdf))

        # PAGE 2
        story.append(PageBreak())

        story.append(Paragraph("<b>Other Bank Account Details</b>", title_style))
        story.append(Spacer(1, 14))

        t2_rows = [
            [
                Paragraph("<b>Sr. No.</b>", body_center),
                Paragraph("<b>Name</b>", body_bold),
                Paragraph("<b>Account No.</b>", body_bold),
                Paragraph("<b>IFSC Code</b>", body_bold),
                Paragraph("<b>Amount</b>", body_right_bold)
            ]
        ]
        for itm in calc_rows:
            t2_rows.append([
                Paragraph(str(itm["sr_no"]), body_center),
                Paragraph(itm["name"], body),
                Paragraph(itm["account_no"], body),
                Paragraph(itm["ifsc_code"], body),
                Paragraph(f"{itm['total']:.2f}", body_right)
            ])

        t2_rows.append([
            "", Paragraph("<b>Total:</b>", body_right_bold), "", "",
            Paragraph(f"<b>Rs. {grand_total:.2f}</b>", body_right_bold)
        ])

        gt2_idx = len(calc_rows) + 1
        t2_pdf = Table(t2_rows, colWidths=[0.6*inch, 2.4*inch, 2.2*inch, 1.1*inch, 1.1*inch])
        t2_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#e2e8f0")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#94a3b8")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('SPAN', (1, gt2_idx), (3, gt2_idx)),
            ('BACKGROUND', (0, gt2_idx), (-1, gt2_idx), colors.HexColor("#f1f5f9")),
        ]))
        story.append(t2_pdf)
        story.append(Spacer(1, 40))

        story.append(Paragraph(
            f"<b>{data.faculty_chairman_title}</b><br/>"
            f"{data.faculty_chairman_name}<br/>"
            f"{data.club_name}<br/><br/>"
            "<font color='#94a3b8'>Signature: ____________________</font>",
            body
        ))

        doc.build(story)
        return output_path
