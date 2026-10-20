import uuid
from html import escape
from pathlib import Path
from typing import Dict, Any, List

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as ReportLabImage, PageBreak, KeepTogether
)

from app.core.config import settings
from app.schemas.submitted_doc import SubmittedDocRequest
from app.services.lead_signature_service import lead_signature_stream
from PIL import Image as PILImage


def _pdf_signature_image(signature_data: str | None):
    signature_stream = lead_signature_stream(signature_data)
    if not signature_stream:
        return None
    width_px, height_px = PILImage.open(signature_stream).size
    signature_stream.seek(0)
    scale = min((1.8 * inch) / width_px, (0.55 * inch) / height_px)
    return ReportLabImage(signature_stream, width=width_px * scale, height=height_px * scale)


def _set_cell_bg(cell, hex_color: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, font_size: float = 9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(15, 23, 42)


def _set_table_borders(table, color="B0B7C3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


class SubmittedDocumentService:

    @staticmethod
    def _add_signature_image(cell, signature_data: str | None):
        signature_stream = lead_signature_stream(signature_data)
        if not signature_stream:
            return
        with PILImage.open(signature_stream) as signature_image:
            ratio = signature_image.width / signature_image.height
            width = Inches(min(1.8, 1.8 if ratio >= 1 else 1.0 * ratio))
            height = Inches(min(0.55, 1.8 / ratio if ratio >= 1 else 1.0))
        signature_stream.seek(0)
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run().add_picture(signature_stream, width=width, height=height)

    @classmethod
    def generate_docx(cls, data: SubmittedDocRequest, output_path: Path):
        doc = Document()
        for sec in doc.sections:
            sec.page_width = Inches(8.27)
            sec.page_height = Inches(11.69)
            sec.top_margin = Inches(0.45)
            sec.bottom_margin = Inches(0.45)
            sec.left_margin = Inches(0.55)
            sec.right_margin = Inches(0.55)

        header_img = settings.PROJECT_ROOT / "assets" / "branding" / "prize_receipt_header.png"
        if header_img.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(str(header_img), width=Inches(7.15))

        # No & Date row
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_before = Pt(4)
        p_meta.paragraph_format.space_after = Pt(6)
        r_no = p_meta.add_run(f"No  :  {data.ref_no}")
        r_no.bold = True
        r_no.font.size = Pt(10)
        p_meta.add_run("\t\t\t\t\t\t\t\t")
        r_date = p_meta.add_run(f"Date  :  {data.date}")
        r_date.bold = True
        r_date.font.size = Pt(10)

        # Submitted to
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(4)
        p_sub.paragraph_format.space_after = Pt(6)
        r_sub = p_sub.add_run(f"Submitted to {data.submitted_to}:")
        r_sub.bold = True
        r_sub.underline = True
        r_sub.font.size = Pt(10)

        # Meta Table
        primary_rows = [
            ("Subject :", data.subject),
            ("Chapter/Club/Society/Section Name :", data.chapter_name),
            ("Chapter/Club/Society/Section Code No :", data.code_no),
            ("Schedule of Event (Date, Day) :", data.event_schedule),
            ("Event Venue :", data.event_venue),
        ]
        t_meta = doc.add_table(rows=len(primary_rows), cols=2)
        t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_meta.autofit = False
        t_meta.columns[0].width = Inches(2.6)
        t_meta.columns[1].width = Inches(4.55)
        for i, (k, v) in enumerate(primary_rows):
            _set_cell_text(t_meta.rows[i].cells[0], k, bold=True)
            _set_cell_text(t_meta.rows[i].cells[1], v)
            _set_cell_bg(t_meta.rows[i].cells[0], "F8FAFC")
        _set_table_borders(t_meta)

        # About the Event Table
        p_sp1 = doc.add_paragraph()
        p_sp1.paragraph_format.space_after = Pt(4)

        about_rows = [
            ("Title of the Event", ":", data.event_title),
            ("Category of the Event", ":", data.category),
            ("Brief Description", ":", data.brief_description),
            ("Objectives", ":", data.objectives),
            ("Target Audience", ":", data.target_audience),
            ("Expected Number of Participants", ":", data.expected_participants),
            ("Name of Judges", ":", data.judges),
        ]
        t_about = doc.add_table(rows=len(about_rows) + 1, cols=3)
        t_about.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_about.autofit = False
        t_about.columns[0].width = Inches(2.4)
        t_about.columns[1].width = Inches(0.25)
        t_about.columns[2].width = Inches(4.5)

        # Section Header Row
        hdr_cell = t_about.rows[0].cells[0]
        hdr_cell.merge(t_about.rows[0].cells[2])
        _set_cell_text(hdr_cell, "About the Event", bold=True, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(hdr_cell, "E2E8F0")

        for idx, (label, sep, val) in enumerate(about_rows):
            row = t_about.rows[idx + 1]
            _set_cell_text(row.cells[0], label, bold=True)
            _set_cell_text(row.cells[1], sep, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(row.cells[2], val)
            _set_cell_bg(row.cells[0], "F8FAFC")
        _set_table_borders(t_about)

        # Financial Proposal Table
        p_sp2 = doc.add_paragraph()
        p_sp2.paragraph_format.space_after = Pt(4)

        fin_items = data.financial_items or []
        t_fin = doc.add_table(rows=len(fin_items) + 3, cols=3)
        t_fin.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_fin.autofit = False
        t_fin.columns[0].width = Inches(1.8)
        t_fin.columns[1].width = Inches(3.35)
        t_fin.columns[2].width = Inches(2.0)

        # Title
        f_hdr = t_fin.rows[0].cells[0]
        f_hdr.merge(t_fin.rows[0].cells[2])
        _set_cell_text(f_hdr, "Financial Proposal", bold=True, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(f_hdr, "E2E8F0")

        # Column Headers
        _set_cell_text(t_fin.rows[1].cells[0], "Event Expenses", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(t_fin.rows[1].cells[1], "Items", bold=True)
        _set_cell_text(t_fin.rows[1].cells[2], "Approximate Amount (Rs)", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        for c in t_fin.rows[1].cells:
            _set_cell_bg(c, "F1F5F9")

        total_amount = sum(float(item.amount) for item in fin_items)
        for i, item in enumerate(fin_items):
            row = t_fin.rows[i + 2]
            if i == 0:
                _set_cell_text(row.cells[0], "")
            else:
                _set_cell_text(row.cells[0], "")
            _set_cell_text(row.cells[1], item.item)
            _set_cell_text(row.cells[2], f"{float(item.amount):.2f}", align=WD_ALIGN_PARAGRAPH.RIGHT)

        # Total Row
        tot_row = t_fin.rows[-1]
        _set_cell_text(tot_row.cells[0], "Total", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(tot_row.cells[1], "")
        _set_cell_text(tot_row.cells[2], f"{total_amount:.2f}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _set_cell_bg(tot_row.cells[0], "F1F5F9")
        _set_cell_bg(tot_row.cells[1], "F1F5F9")
        _set_cell_bg(tot_row.cells[2], "F1F5F9")
        _set_table_borders(t_fin)

        # PAGE 2
        doc.add_page_break()

        # Signatures Grid (2x2)
        p_sig_hdr = doc.add_paragraph()
        p_sig_hdr.paragraph_format.space_before = Pt(8)
        p_sig_hdr.paragraph_format.space_after = Pt(6)

        t_sig = doc.add_table(rows=2, cols=2)
        t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sig.autofit = False
        t_sig.columns[0].width = Inches(3.55)
        t_sig.columns[1].width = Inches(3.55)

        sig_cell_00 = (
            f"Signature of Student Head\n"
            f"Name: {data.student_head_name}\n"
            f"Roll No. {data.student_head_roll}\n"
            f"Branch: {data.student_head_branch}\n"
            f"Mobile No. {data.student_head_mobile}"
        )
        sig_cell_01 = (
            f"Signature of Student Secretary\n"
            f"Name: {data.student_secretary_name}\n"
            f"Roll No. {data.student_secretary_roll}\n"
            f"Branch: {data.student_secretary_branch}\n"
            f"Mobile No. {data.student_secretary_mobile}"
        )
        sig_cell_10 = (
            f"Signature of Faculty Co-Chairman\n"
            f"Name: {data.faculty_co_chairman_name}\n"
            f"Designation: {data.faculty_co_chairman_designation}\n"
            f"Dept: {data.faculty_co_chairman_dept}\n"
            f"Mobile No. {data.faculty_co_chairman_mobile}"
        )
        sig_cell_11 = (
            f"Signature of Faculty Chairman\n"
            f"Name: {data.faculty_chairman_name}\n"
            f"Designation: {data.faculty_chairman_designation}\n"
            f"Dept: {data.faculty_chairman_dept}\n"
            f"Mobile No. {data.faculty_chairman_mobile}"
        )
        _set_cell_text(t_sig.rows[0].cells[0], sig_cell_00)
        _set_cell_text(t_sig.rows[0].cells[1], sig_cell_01)
        cls._add_signature_image(t_sig.rows[0].cells[0], data.student_head_signature_data)
        cls._add_signature_image(t_sig.rows[0].cells[1], data.student_secretary_signature_data)
        _set_cell_text(t_sig.rows[1].cells[0], sig_cell_10)
        _set_cell_text(t_sig.rows[1].cells[1], sig_cell_11)

        # Office Use Table
        p_sp3 = doc.add_paragraph()
        p_sp3.paragraph_format.space_before = Pt(14)
        p_sp3.paragraph_format.space_after = Pt(4)

        office_fields = [
            "Total amount Sanctioned in Current Academic Year :",
            "Amount Utilized :",
            "Current Proposed amount :",
            "Balance amount :",
            "Signature (Office Clerk) :",
        ]
        t_off = doc.add_table(rows=len(office_fields) + 1, cols=2)
        t_off.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_off.autofit = False
        t_off.columns[0].width = Inches(5.15)
        t_off.columns[1].width = Inches(2.0)

        hdr_off = t_off.rows[0].cells[0]
        hdr_off.merge(t_off.rows[0].cells[1])
        _set_cell_text(hdr_off, "Office Use", bold=True, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(hdr_off, "E2E8F0")

        for idx, lbl in enumerate(office_fields):
            row = t_off.rows[idx + 1]
            _set_cell_text(row.cells[0], lbl, bold=True)
            _set_cell_text(row.cells[1], "")
        _set_table_borders(t_off)

        # Authority Approvals (3 columns)
        p_sp4 = doc.add_paragraph()
        p_sp4.paragraph_format.space_before = Pt(24)
        p_sp4.paragraph_format.space_after = Pt(4)

        t_auth = doc.add_table(rows=3, cols=3)
        t_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_auth.autofit = False
        for c in t_auth.columns:
            c.width = Inches(2.38)

        _set_cell_text(t_auth.rows[0].cells[0], "Recommended / Not Recommended", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(t_auth.rows[0].cells[1], "Recommended / Not Recommended", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(t_auth.rows[0].cells[2], "Approved / Not Approved", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        _set_cell_text(t_auth.rows[1].cells[0], "\n\n")
        _set_cell_text(t_auth.rows[1].cells[1], "\n\n")
        _set_cell_text(t_auth.rows[1].cells[2], "\n\n")

        _set_cell_text(t_auth.rows[2].cells[0], f"{data.chairman_fsw}\nChairman, F& SW", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(t_auth.rows[2].cells[1], f"{data.associate_dean_sw}\nAssociate Dean(SW)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(t_auth.rows[2].cells[2], f"{data.dean_sw}\nDean (Students Welfare)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        doc.save(output_path)

    @classmethod
    def generate_pdf(cls, data: SubmittedDocRequest, output_path: Path):
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=14 * mm,
            rightMargin=14 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()

        normal = ParagraphStyle("SubNormal", parent=styles["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#0f172a"))
        normal_b = ParagraphStyle("SubNormalBold", parent=normal, fontName="Helvetica-Bold")
        center_b = ParagraphStyle("SubCenterBold", parent=normal_b, alignment=1)
        right_b = ParagraphStyle("SubRightBold", parent=normal_b, alignment=2)
        right_norm = ParagraphStyle("SubRightNorm", parent=normal, alignment=2)

        story = []

        # 1. Trilingual Header Image
        header_img = settings.PROJECT_ROOT / "assets" / "branding" / "prize_receipt_header.png"
        if header_img.exists():
            story.append(ReportLabImage(str(header_img), width=182 * mm, height=32.47 * mm))
            story.append(Spacer(1, 2 * mm))

        # 2. No and Date Row
        no_date_data = [
            [
                Paragraph(f"<b>No : {escape(data.ref_no)}</b>", normal),
                Paragraph(f"<b>Date : {escape(data.date)}</b>", right_b),
            ]
        ]
        t_no_date = Table(no_date_data, colWidths=[100 * mm, 82 * mm], hAlign="CENTER")
        t_no_date.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(t_no_date)
        story.append(Spacer(1, 2 * mm))

        # 3. Submitted to
        story.append(Paragraph(f"<b><u>Submitted to {escape(data.submitted_to)}:</u></b>", ParagraphStyle("SubTo", parent=normal_b, fontSize=9.5, leading=13)))
        story.append(Spacer(1, 2 * mm))

        # 4. Meta Table
        primary_rows = [
            [Paragraph("<b>Subject :</b>", normal), Paragraph(escape(data.subject), normal)],
            [Paragraph("<b>Chapter/Club/Society/Section Name :</b>", normal), Paragraph(escape(data.chapter_name), normal)],
            [Paragraph("<b>Chapter/Club/Society/Section Code No :</b>", normal), Paragraph(escape(data.code_no), normal)],
            [Paragraph("<b>Schedule of Event (Date, Day) :</b>", normal), Paragraph(escape(data.event_schedule), normal)],
            [Paragraph("<b>Event Venue :</b>", normal), Paragraph(escape(data.event_venue), normal)],
        ]
        t_meta = Table(primary_rows, colWidths=[65 * mm, 117 * mm], hAlign="CENTER")
        t_meta.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#334155")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t_meta)
        story.append(Spacer(1, 3 * mm))

        # 5. About the Event Table
        about_data = [
            [Paragraph("<b>About the Event</b>", center_b), "", ""],
            [Paragraph("<b>Title of the Event</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.event_title), normal)],
            [Paragraph("<b>Category of the Event</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.category), normal)],
            [Paragraph("<b>Brief Description</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.brief_description), normal)],
            [Paragraph("<b>Objectives</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.objectives), normal)],
            [Paragraph("<b>Target Audience</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.target_audience), normal)],
            [Paragraph("<b>Expected Number of Participants</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.expected_participants), normal)],
            [Paragraph("<b>Name of Judges</b>", normal), Paragraph(":", center_b), Paragraph(escape(data.judges), normal)],
        ]
        t_about = Table(about_data, colWidths=[55 * mm, 6 * mm, 121 * mm], hAlign="CENTER")
        t_about.setStyle(TableStyle([
            ("SPAN", (0, 0), (2, 0)),
            ("BACKGROUND", (0, 0), (2, 0), colors.HexColor("#e2e8f0")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#334155")),
            ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#f8fafc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t_about)
        story.append(Spacer(1, 3 * mm))

        # 6. Financial Proposal Table
        fin_items = data.financial_items or []
        total_amount = sum(float(item.amount) for item in fin_items)
        fin_data = [
            [Paragraph("<b>Financial Proposal</b>", center_b), "", ""],
            [Paragraph("<b>Event Expenses</b>", center_b), Paragraph("<b>Items</b>", normal_b), Paragraph("<b>Approximate Amount (Rs)</b>", right_b)],
        ]
        for i, item in enumerate(fin_items):
            fin_data.append([
                Paragraph("", normal),
                Paragraph(escape(item.item), normal),
                Paragraph(f"{float(item.amount):.2f}", right_norm),
            ])
        fin_data.append([
            Paragraph("<b>Total</b>", center_b),
            Paragraph("", normal),
            Paragraph(f"<b>{total_amount:.2f}</b>", right_b),
        ])

        t_fin = Table(fin_data, colWidths=[45 * mm, 87 * mm, 50 * mm], hAlign="CENTER")
        t_fin.setStyle(TableStyle([
            ("SPAN", (0, 0), (2, 0)),
            ("BACKGROUND", (0, 0), (2, 0), colors.HexColor("#e2e8f0")),
            ("BACKGROUND", (0, 1), (2, 1), colors.HexColor("#f1f5f9")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#334155")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f1f5f9")),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(t_fin)

        # PAGE 2
        story.append(PageBreak())

        # Signatures 2x2
        sig_data = [
            [
                Paragraph(
                    f"<b>Signature of Student Head</b><br/>"
                    f"Name: {escape(data.student_head_name)}<br/>"
                    f"Roll No. {escape(data.student_head_roll)}<br/>"
                    f"Branch: {escape(data.student_head_branch)}<br/>"
                    f"Mobile No. {escape(data.student_head_mobile)}",
                    normal
                ) if not data.student_head_signature_data else [
                    Paragraph(f"<b>Signature of Student Head</b><br/>Name: {escape(data.student_head_name)}<br/>Roll No. {escape(data.student_head_roll)}<br/>Branch: {escape(data.student_head_branch)}<br/>Mobile No. {escape(data.student_head_mobile)}", normal),
                    _pdf_signature_image(data.student_head_signature_data),
                ],
                Paragraph(
                    f"<b>Signature of Student Secretary</b><br/>"
                    f"Name: {escape(data.student_secretary_name)}<br/>"
                    f"Roll No. {escape(data.student_secretary_roll)}<br/>"
                    f"Branch: {escape(data.student_secretary_branch)}<br/>"
                    f"Mobile No. {escape(data.student_secretary_mobile)}",
                    normal
                ) if not data.student_secretary_signature_data else [
                    Paragraph(f"<b>Signature of Student Secretary</b><br/>Name: {escape(data.student_secretary_name)}<br/>Roll No. {escape(data.student_secretary_roll)}<br/>Branch: {escape(data.student_secretary_branch)}<br/>Mobile No. {escape(data.student_secretary_mobile)}", normal),
                    _pdf_signature_image(data.student_secretary_signature_data),
                ],
            ],
            [
                Paragraph(
                    f"<b>Signature of Faculty Co-Chairman</b><br/>"
                    f"Name: {escape(data.faculty_co_chairman_name)}<br/>"
                    f"Designation: {escape(data.faculty_co_chairman_designation)}<br/>"
                    f"Dept: {escape(data.faculty_co_chairman_dept)}<br/>"
                    f"Mobile No. {escape(data.faculty_co_chairman_mobile)}",
                    normal
                ),
                Paragraph(
                    f"<b>Signature of Faculty Chairman</b><br/>"
                    f"Name: {escape(data.faculty_chairman_name)}<br/>"
                    f"Designation: {escape(data.faculty_chairman_designation)}<br/>"
                    f"Dept: {escape(data.faculty_chairman_dept)}<br/>"
                    f"Mobile No. {escape(data.faculty_chairman_mobile)}",
                    normal
                ),
            ],
        ]
        t_sig = Table(sig_data, colWidths=[91 * mm, 91 * mm], hAlign="CENTER")
        t_sig.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(t_sig)
        story.append(Spacer(1, 8 * mm))

        # Office Use Table
        office_rows = [
            [Paragraph("<b>Office Use</b>", center_b), ""],
            [Paragraph("<b>Total amount Sanctioned in Current Academic Year :</b>", normal), Paragraph(escape(data.office_sanctioned_amount or ""), normal)],
            [Paragraph("<b>Amount Utilized :</b>", normal), Paragraph(escape(data.office_amount_utilized or ""), normal)],
            [Paragraph("<b>Current Proposed amount :</b>", normal), Paragraph(escape(data.office_proposed_amount or ""), normal)],
            [Paragraph("<b>Balance amount :</b>", normal), Paragraph(escape(data.office_balance_amount or ""), normal)],
            [Paragraph("<b>Signature (Office Clerk) :</b>", normal), Paragraph(escape(data.office_clerk_signature or ""), normal)],
        ]
        t_off = Table(office_rows, colWidths=[120 * mm, 62 * mm], hAlign="CENTER")
        t_off.setStyle(TableStyle([
            ("SPAN", (0, 0), (1, 0)),
            ("BACKGROUND", (0, 0), (1, 0), colors.HexColor("#e2e8f0")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#334155")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_off)
        story.append(Spacer(1, 14 * mm))

        # Dean SW 3 Approvers
        auth_data = [
            [
                Paragraph("<b>Recommended / Not Recommended</b>", center_b),
                Paragraph("<b>Recommended / Not Recommended</b>", center_b),
                Paragraph("<b>Approved / Not Approved</b>", center_b),
            ],
            [Paragraph("<br/><br/>", normal), Paragraph("<br/><br/>", normal), Paragraph("<br/><br/>", normal)],
            [
                Paragraph(f"<b>{escape(data.chairman_fsw)}</b><br/>Chairman, F& SW", center_b),
                Paragraph(f"<b>{escape(data.associate_dean_sw)}</b><br/>Associate Dean(SW)", center_b),
                Paragraph(f"<b>{escape(data.dean_sw)}</b><br/>Dean (Students Welfare)", center_b),
            ],
        ]
        t_auth = Table(auth_data, colWidths=[60 * mm, 62 * mm, 60 * mm], hAlign="CENTER")
        t_auth.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t_auth)

        doc.build(story)
