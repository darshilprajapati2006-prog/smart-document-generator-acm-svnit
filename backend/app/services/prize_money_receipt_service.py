from decimal import Decimal
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings
from app.schemas.prize_money_receipt import PrizeMoneyReceiptRequest


def _money(value: Decimal) -> str:
    return f"{value:,.2f}"


def _shade(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_borders(table, color: str = "8C8C8C") -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), color)
        borders.append(node)
    table._tbl.tblPr.append(borders)


def _set_cell_text(cell, value: str, *, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)


def _header_values(data: PrizeMoneyReceiptRequest):
    return [
        ("Serial Number", data.serial_number),
        ("Date", data.date),
    ]


def _acknowledgement(data: PrizeMoneyReceiptRequest) -> str:
    return (
        f"It is hereby acknowledged that an amount of ₹ {_money(data.amount_rupees)} will be received by me "
        f"for securing {data.position} in the competition/event named {data.competition_name}, "
        f"organized by {data.organizing_body}, held on {data.event_date} at {data.venue}."
    )


def _account_values(data: PrizeMoneyReceiptRequest):
    return [
        ("Affiliating Institute", data.affiliating_institute),
        ("Name as per Bank Account", data.name_as_per_bank),
        ("Bank Name", data.bank_name),
        ("Branch", data.branch),
        ("Account Number", data.account_number),
        ("IFSC Code", data.ifsc_code),
        ("PAN Card Number", data.pan_number),
    ]


def _student_values(data: PrizeMoneyReceiptRequest):
    return [
        ("Student Name", data.student_name),
        ("Roll Number", data.roll_number),
        ("Course and Semester", data.course_semester),
        ("Mobile Number", data.mobile_number),
    ]


class PrizeMoneyReceiptService:
    @classmethod
    def generate_docx(cls, data: PrizeMoneyReceiptRequest, output_path: Path) -> None:
        doc = Document()
        section = doc.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

        header_path = settings.PROJECT_ROOT / "assets" / "branding" / "prize_receipt_header.png"
        header_paragraph = doc.add_paragraph()
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_paragraph.paragraph_format.space_after = Pt(3)
        header_paragraph.add_run().add_picture(str(header_path), width=Inches(7.15))

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("PRIZE MONEY RECEIPT")
        run.bold = True
        run.underline = True
        run.font.size = Pt(13)

        table = doc.add_table(rows=len(_header_values(data)), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(2.15)
        table.columns[1].width = Inches(5.15)
        for row, (label, value) in zip(table.rows, _header_values(data)):
            row.cells[0].width = Inches(2.15)
            row.cells[1].width = Inches(5.15)
            _shade(row.cells[0], "F1F3F5")
            _set_cell_text(row.cells[0], label, bold=True)
            _set_cell_text(row.cells[1], value)
        _set_borders(table)

        p = doc.add_paragraph(_acknowledgement(data))
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.08
        p.runs[0].font.size = Pt(9)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("(Account Detail)")
        r.bold = True
        r.font.size = Pt(10)
        info = doc.add_table(rows=len(_account_values(data)), cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        info.autofit = False
        info.columns[0].width = Inches(3.0)
        info.columns[1].width = Inches(4.3)
        for row, (label, value) in zip(info.rows, _account_values(data)):
            row.cells[0].width = Inches(3.0)
            row.cells[1].width = Inches(4.3)
            _shade(row.cells[0], "F1F3F5")
            _set_cell_text(row.cells[0], f"{label}:", bold=True)
            _set_cell_text(row.cells[1], value)
        _set_borders(info)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.add_run("Student Details").bold = True
        for label, value in _student_values(data):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            p.add_run(f"{label}: {value}").font.size = Pt(9)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        p.add_run("Signature: ____________________").font.size = Pt(9)
        p = doc.add_paragraph("It is certified that the said amount is recommended for payment.")
        p.paragraph_format.space_after = Pt(15)
        p.runs[0].font.size = Pt(9)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        chairman = p.add_run(f"Chairman Signature: ____________________    {data.chairman_name}")
        chairman.font.size = Pt(9)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

    @classmethod
    def generate_pdf(cls, data: PrizeMoneyReceiptRequest, output_path: Path) -> None:
        font_path = Path(__file__).resolve().parents[1] / "assets" / "fonts" / "DejaVuSans.ttf"
        if "SmartDocsUnicode" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("SmartDocsUnicode", str(font_path)))

        doc = SimpleDocTemplate(
            str(output_path), pagesize=A4,
            rightMargin=15 * mm, leftMargin=15 * mm,
            topMargin=10 * mm, bottomMargin=12 * mm,
            title="Prize Money Receipt - ACM SVNIT",
        )
        normal = ParagraphStyle("ReceiptText", fontName="SmartDocsUnicode", fontSize=8.5, leading=11, textColor=colors.HexColor("#1f2937"))
        center = ParagraphStyle("ReceiptCenter", parent=normal, alignment=TA_CENTER)
        section_style = ParagraphStyle("ReceiptSection", parent=normal, fontSize=9, leading=12, spaceBefore=5, spaceAfter=4)
        story = []
        header_path = settings.PROJECT_ROOT / "assets" / "branding" / "prize_receipt_header.png"
        story.extend([
            Image(str(header_path), width=180 * mm, height=32.47 * mm),
            Spacer(1, 2 * mm),
            Paragraph("<b><u>Prize Money Receipt</u></b>", ParagraphStyle("ReceiptTitle", parent=center, fontSize=12, leading=15)),
            Spacer(1, 2 * mm),
        ])

        def pairs_table(pairs):
            values = [[Paragraph(f"<b>{escape(label)}</b>", normal), Paragraph(escape(value).replace("\n", "<br/>"), normal)] for label, value in pairs]
            table = Table(values, colWidths=[52 * mm, 137 * mm], hAlign="CENTER", splitByRow=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.55, colors.HexColor("#888888")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f3f5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            return table

        story.append(pairs_table(_header_values(data)))
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(escape(_acknowledgement(data)), normal))
        story.append(Paragraph("<b>(Account Detail)</b>", ParagraphStyle("AccountTitle", parent=center, spaceBefore=4, spaceAfter=4)))
        story.append(pairs_table(_account_values(data)))
        story.append(Paragraph("<b>Student Details</b>", section_style))
        story.extend(Paragraph(f"•&nbsp;&nbsp; <b>{escape(label)}:</b> {escape(value)}", normal) for label, value in _student_values(data))
        story.extend([
            Spacer(1, 2 * mm), Paragraph("Signature: ____________________", normal),
            Spacer(1, 4 * mm), Paragraph("It is certified that the said amount is recommended for payment.", normal),
            Spacer(1, 8 * mm), Paragraph(f"Chairman Signature: ____________________&nbsp;&nbsp; {escape(data.chairman_name)}", normal),
        ])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        def paint_white_page(canvas, _document):
            canvas.saveState()
            canvas.setFillColor(colors.white)
            canvas.rect(0, 0, A4[0], A4[1], stroke=0, fill=1)
            canvas.restoreState()

        doc.build(story, onFirstPage=paint_white_page, onLaterPages=paint_white_page)
