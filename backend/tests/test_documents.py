import pytest
from pathlib import Path
from fastapi.testclient import TestClient
from app.main import app
from app.services.executive_meet_service import ExecutiveMeetService
from app.services.bill_certificate_service import BillCertificateService
from app.services.bill_summary_service import BillSummaryService
from app.schemas.executive_meet import ExecutiveMeetRequest
from app.schemas.bill_certificate import BillCertificateRequest, BillItem
from app.schemas.bill_summary import BillSummaryRequest, BillSummaryRow

client = TestClient(app)

def test_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert "Smart Document Generator API" in data["message"]

def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"

def test_list_templates():
    response = client.get("/api/documents/templates")
    assert response.status_code == 200
    data = response.json()
    assert "templates" in data
    ids = [t["id"] for t in data["templates"]]
    assert "exec_meet" in ids
    assert "bill_certificate" in ids
    assert "bill_summary" in ids

def test_executive_meet_generation(tmp_path):
    payload = {
        "ref_no": "ACM/26-27/ROOM/009",
        "date": "09/10/2026",
        "submitted_to": "HoD(CSE Dept.)",
        "subject": "Requesting Permission to Conduct ACM Executives Meet in Seminar Hall 402, CSE Department.",
        "event_name": "ACM Executives Meet",
        "venue": "Seminar Hall 402, CSE Department",
        "event_date": "15/10/2026",
        "time_from": "06:00 PM",
        "time_to": "07:30 PM",
        "brief_description": "ACM NIT Surat is organizing an Executive Meet to discuss the chapter vision and activities.",
        "objective": "To provide guidance to the ACM executive team on organizing events.",
        "participants": "150+",
        "mode": "Offline",
        "organizers": "ACM Core 2026-27"
    }
    response = client.post("/api/documents/executive-meet/generate", json=payload)
    assert response.status_code == 200
    res_data = response.json()
    assert res_data["success"] is True
    assert "document_id" in res_data
    assert "files" in res_data
    assert "docx" in res_data["files"]
    assert "pdf" in res_data["files"]

    doc_id = res_data["document_id"]
    # Download DOCX
    dl_docx = client.get(f"/api/documents/download/{doc_id}/docx")
    assert dl_docx.status_code == 200
    assert dl_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(dl_docx.content) > 1000

    # Download PDF
    dl_pdf = client.get(f"/api/documents/download/{doc_id}/pdf")
    assert dl_pdf.status_code == 200
    assert dl_pdf.headers["content-type"] == "application/pdf"
    assert len(dl_pdf.content) > 1000

def test_executive_meet_validation():
    # Missing required fields like date, event_date, brief_description
    response = client.post("/api/documents/executive-meet/generate", json={})
    assert response.status_code == 422

def test_bill_certificate_calculation_and_generation():
    payload = {
        "chapter_name": "ACM",
        "code_no": "6/52",
        "approval_note_no": "ACM/2026-27/SIH/001",
        "date": "08/09/2026",
        "supplier_name": "Stationery & Tech Supplies Ltd.",
        "items": [
            {"name": "Item Alpha", "quantity": 2, "unit_cost": 250.0},
            {"name": "Item Beta", "quantity": 5, "unit_cost": 100.0}
        ],
        "taxes_or_other_charges": 50.0,
        "invoice_no_and_date": "INV-1029 dtd 08/09/2026"
    }
    response = client.post("/api/documents/bill-certificate/generate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    summary = data["calculation_summary"]
    assert summary["subtotal"] == 1000.0
    assert summary["taxes"] == 50.0
    assert summary["grand_total"] == 1050.0

    doc_id = data["document_id"]
    dl_docx = client.get(f"/api/documents/download/{doc_id}/docx")
    assert dl_docx.status_code == 200
    dl_pdf = client.get(f"/api/documents/download/{doc_id}/pdf")
    assert dl_pdf.status_code == 200

def test_bill_certificate_invalid_items():
    payload = {
        "chapter_name": "ACM",
        "date": "08/09/2026",
        "supplier_name": "Test Vendor",
        "items": [
            {"name": "Bad Item", "quantity": 0, "unit_cost": 100.0}
        ]
    }
    response = client.post("/api/documents/bill-certificate/generate", json=payload)
    assert response.status_code == 422

def test_bill_summary_multiple_recipients():
    payload = {
        "date": "09/10/2026",
        "title": "BILL SUMMARY",
        "subtitle": "SIH Internal Hackathon Prize Distributions",
        "rows": [
            {
                "name": "Team Codex",
                "account_details": "A/C: 1122334455, IFSC: SBIN0001234",
                "to_be_paid_to": "Ansh Gupta",
                "prize_money": 5000.0,
                "other_amount": 250.0,
                "account_no": "1122334455",
                "ifsc_code": "SBIN0001234"
            },
            {
                "name": "Team Algo",
                "account_details": "A/C: 9988776655, IFSC: HDFC0005678",
                "to_be_paid_to": "Arshad Khatib",
                "prize_money": 3000.0,
                "other_amount": 0.0,
                "account_no": "9988776655",
                "ifsc_code": "HDFC0005678"
            }
        ]
    }
    response = client.post("/api/documents/bill-summary/generate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    summary = data["calculation_summary"]
    assert summary["total_prize_money"] == 8000.0
    assert summary["total_other"] == 250.0
    assert summary["grand_total"] == 8250.0

    doc_id = data["document_id"]
    dl_docx = client.get(f"/api/documents/download/{doc_id}/docx")
    assert dl_docx.status_code == 200
    dl_pdf = client.get(f"/api/documents/download/{doc_id}/pdf")
    assert dl_pdf.status_code == 200

def test_invalid_download_id_and_format():
    response = client.get("/api/documents/download/non-existent-uuid/docx")
    assert response.status_code in [400, 404]

    response = client.get("/api/documents/download/12345678-1234-1234-1234-123456789abc/exe")
    assert response.status_code in [400, 404]

# TEST B — Executive Meet partial extraction
def test_nlp_test_b_partial_extraction():
    prompt = "Organize an AI workshop at Seminar Hall 402 for 120 students."
    response = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert response.status_code == 200
    data = response.json()
    fields = data["extracted_fields"]
    assert "event_name" in fields
    assert "AI Workshop" in fields["event_name"]
    assert fields.get("venue") == "Seminar Hall 402"
    assert fields.get("participants") == "120"
    # Unmentioned fields must NOT be fabricated
    assert "event_date" not in fields
    assert "time_from" not in fields
    assert "time_to" not in fields
    assert "ref_no" not in fields
    assert "lead_name" not in fields

# TEST C — Executive Meet complete extraction
def test_nlp_test_c_complete_extraction():
    prompt = (
        "Organize an ACM AI and Machine Learning workshop on 20 November 2026 from 10 AM to 1 PM "
        "at Seminar Hall 402, SVNIT Surat, for 120 students. The workshop will be organized by the ACM Core Team."
    )
    response = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert response.status_code == 200
    data = response.json()
    fields = data["extracted_fields"]
    assert "AI and Machine Learning" in fields["event_name"]
    assert fields["event_date"] in ["20 November 2026", "20/11/2026"]
    assert fields["time_from"] == "10:00 AM"
    assert fields["time_to"] in ["01:00 PM", "1:00 PM"]
    assert fields["venue"] == "Seminar Hall 402, SVNIT Surat"
    assert fields["participants"] == "120"
    assert fields["organizers"] == "ACM Core Team"

# Complete ML Workshop Example Test (Item 5 & 10)
def test_nlp_ml_workshop_complete_prompt():
    prompt = (
        "Create an executive meeting document for a Machine Learning Workshop organized by ACM Student Chapter, SVNIT Surat.\n\n"
        "The workshop is titled 'Introduction to Machine Learning and Generative AI' and will be conducted on 28 October 2026 from 10:00 AM to 4:00 PM at the Seminar Hall, SVNIT Surat.\n\n"
        "The workshop will be coordinated by Darshil Prajapati. The target audience is first-year and second-year engineering students. The session will cover machine learning fundamentals, neural networks, practical Python demonstrations, and an introduction to generative AI.\n\n"
        "The estimated budget is ₹15,000, including ₹6,000 for refreshments, ₹4,000 for workshop materials, ₹3,000 for banners and printing, and ₹2,000 for miscellaneous expenses.\n\n"
        "The expected number of participants is 100. Registration will be free for all SVNIT students."
    )
    response = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert response.status_code == 200
    data = response.json()
    assert data["recommended_type"] == "exec_meet"
    fields = data["extracted_fields"]

    assert fields["event_name"] == "Introduction to Machine Learning and Generative AI"
    assert fields["event_date"] == "28 October 2026"
    assert fields["venue"] == "Seminar Hall, SVNIT Surat"
    assert fields["time_from"] == "10:00 AM"
    assert fields["time_to"] == "4:00 PM"
    assert fields["participants"] == "100"
    assert fields["organizers"] == "ACM Student Chapter, SVNIT Surat"
    assert fields["lead_name"] == "Darshil Prajapati"
    assert "A workshop covering machine learning fundamentals, neural networks, practical Python demonstrations" in fields["brief_description"]

    # Verify alias mappings are provided
    assert fields["start_time"] == "10:00 AM"
    assert fields["end_time"] == "4:00 PM"
    assert fields["expected_participants"] == "100"
    assert fields["organizing_club"] == "ACM Student Chapter, SVNIT Surat"
    assert fields["coordinator"] == "Darshil Prajapati"

    # Verify no unmentioned administrative fields fabricated
    assert "mode" not in fields
    assert "ref_no" not in fields
    assert "submitted_to" not in fields
    assert "subject" not in fields

# Missing information test (Item 10)
def test_nlp_missing_information_partial():
    prompt = "Conduct an Executive Meet at Seminar Hall 402 on 15 November 2026."
    response = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert response.status_code == 200
    fields = response.json()["extracted_fields"]
    assert "Executive Meet" in fields["event_name"]
    assert fields["venue"] == "Seminar Hall 402"
    assert fields["event_date"] == "15 November 2026"
    # Unmentioned fields must remain missing
    assert "time_from" not in fields
    assert "time_to" not in fields
    assert "participants" not in fields
    assert "organizers" not in fields
    assert "brief_description" not in fields
    assert "lead_name" not in fields

# TEST D — No fabrication
def test_nlp_test_d_no_fabrication():
    prompt = "Organize a workshop."
    response = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert response.status_code == 200
    fields = response.json()["extracted_fields"]
    assert fields["event_name"] == "Workshop"
    # Crucial: No fabricated dates, venue, counts, reference numbers, or signatories
    assert "event_date" not in fields
    assert "venue" not in fields
    assert "participants" not in fields
    assert "ref_no" not in fields
    assert "lead_name" not in fields
    assert "objective" not in fields

# TEST E — Bill Certificate extraction
def test_nlp_test_e_bill_certificate():
    prompt = "Purchase 25 notebooks at Rs 80 each and 10 pen sets at Rs 120 each from ABC Stationery."
    response = client.post("/api/documents/nlp/extract", json={"text": prompt, "document_type": "bill_certificate"})
    assert response.status_code == 200
    data = response.json()
    assert data["recommended_type"] == "bill_certificate"
    items = data["items"]
    assert len(items) == 2
    assert items[0]["name"] == "Notebooks"
    assert items[0]["quantity"] == 25
    assert items[0]["unit_cost"] == 80.0
    assert items[0]["total_cost"] == 2000.0

    assert items[1]["name"] == "Pen Sets"
    assert items[1]["quantity"] == 10
    assert items[1]["unit_cost"] == 120.0
    assert items[1]["total_cost"] == 1200.0

    assert data["extracted_fields"]["supplier_name"] == "ABC Stationery"
    assert data["calculation_summary"]["subtotal"] == 3200.0
    # No fabricated bank or GST details
    assert "account_number" not in data["extracted_fields"]
    assert "supplier_gst" not in data["extracted_fields"]

# TEST F — Bill Summary extraction
def test_nlp_test_f_bill_summary():
    prompt = "Prepare a reimbursement summary for three students: Rahul Rs 1500 for travel, Priya Rs 800 for printing, and Amit Rs 1200 for event supplies."
    response = client.post("/api/documents/nlp/extract", json={"text": prompt, "document_type": "bill_summary"})
    assert response.status_code == 200
    data = response.json()
    assert data["recommended_type"] == "bill_summary"
    recipients = data["recipients"]
    assert len(recipients) == 3
    assert recipients[0]["name"] == "Rahul"
    assert recipients[0]["prize_money"] == 1500.0
    assert recipients[1]["name"] == "Priya"
    assert recipients[1]["prize_money"] == 800.0
    assert recipients[2]["name"] == "Amit"
    assert recipients[2]["prize_money"] == 1200.0
    assert data["calculation_summary"]["grand_total"] == 3500.0
    # No fabricated bank accounts or signatures
    assert not recipients[0]["account_no"]
    assert "student_head_name" not in data["extracted_fields"]

# TEST G & H — State isolation & empty prompt
def test_nlp_test_empty_prompt():
    response = client.post("/api/documents/nlp/extract", json={"text": ""})
    assert response.status_code == 200
    data = response.json()
    assert data["extracted_fields"] == {}
    assert data["items"] == []
    assert data["recipients"] == []

def test_long_text_and_multipage():
    long_desc = "Lorem ipsum dolor sit amet, consectetur adipiscing elit. " * 30
    long_obj = "Comprehensive practical hands-on workshop objective. " * 20
    payload = {
        "ref_no": "ACM/26-27/ROOM/099",
        "date": "09/10/2026",
        "submitted_to": "HoD(CSE Dept.)",
        "subject": "Requesting Permission for Extended Hackathon in Lab 402.",
        "event_name": "ACM Extended Hackathon 2026",
        "venue": "Seminar Hall 402 & CSE Lab 1",
        "event_date": "20/10/2026",
        "time_from": "09:00 AM",
        "time_to": "09:00 PM",
        "brief_description": long_desc,
        "objective": long_obj,
        "participants": "250+",
        "mode": "Hybrid",
        "organizers": "ACM SVNIT Chapter"
    }
    response = client.post("/api/documents/executive-meet/generate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True

    doc_id = data["document_id"]
    r_docx = client.get(f"/api/documents/download/{doc_id}/docx")
    r_pdf = client.get(f"/api/documents/download/{doc_id}/pdf")
    assert r_docx.status_code == 200
    assert r_pdf.status_code == 200

def test_decimal_financial_accuracy():
    payload = {
        "chapter_name": "ACM",
        "code_no": "6/52",
        "approval_note_no": "ACM/2026-27/NOTE/005",
        "date": "09/10/2026",
        "supplier_name": "Precision Tech Components",
        "items": [
            {"name": "Electronic Sensors", "quantity": 3, "unit_cost": 33.33}
        ],
        "taxes_or_other_charges": 18.01
    }
    response = client.post("/api/documents/bill-certificate/generate", json=payload)
    assert response.status_code == 200
    summary = response.json()["calculation_summary"]
    assert summary["subtotal"] == 99.99
    assert summary["taxes"] == 18.01
    assert summary["grand_total"] == 118.00

def test_get_template_details_valid_and_invalid():
    # Valid template details
    res_valid = client.get("/api/documents/templates/exec_meet")
    assert res_valid.status_code == 200
    data = res_valid.json()
    assert data["id"] == "exec_meet"
    assert "required_fields" in data
    assert "event_name" in data["required_fields"]
    assert "field_labels" in data

    # Invalid template ID
    res_invalid = client.get("/api/documents/templates/unknown_template_xyz")
    assert res_invalid.status_code == 404

def test_missing_required_fields_reported_in_nlp():
    prompt = "Organize a seminar at Seminar Hall on 12/12/2026"
    res = client.post("/api/documents/nlp/extract", json={"text": prompt})
    assert res.status_code == 200
    data = res.json()
    assert "missing_required_fields" in data
    missing = data["missing_required_fields"]
    assert isinstance(missing, list)
    # Start Time, End Time, Brief Description should be flagged as missing
    assert any("Time" in m for m in missing)
    assert any("Brief Description" in m for m in missing)

def test_docx_and_pdf_binary_validity_and_content():
    import io
    from docx import Document
    
    unique_event_name = "Super Quantum Computing Seminar 2026"
    payload = {
        "ref_no": "ACM/26-27/ROOM/999",
        "date": "10/10/2026",
        "submitted_to": "HoD(CSE Dept.)",
        "subject": f"Requesting Permission for {unique_event_name}",
        "event_name": unique_event_name,
        "venue": "Seminar Hall 402, SVNIT Surat",
        "event_date": "18/10/2026",
        "time_from": "02:00 PM",
        "time_to": "05:00 PM",
        "brief_description": "Exclusive quantum computing practical demonstrations.",
        "objective": "To introduce quantum logic gates.",
        "participants": "100",
        "mode": "Offline",
        "organizers": "ACM Student Chapter, SVNIT Surat",
        "lead_name": "Darshil Prajapati"
    }
    res = client.post("/api/documents/executive-meet/generate", json=payload)
    assert res.status_code == 200
    doc_id = res.json()["document_id"]

    # Verify DOCX binary structure and text content
    r_docx = client.get(f"/api/documents/download/{doc_id}/docx")
    assert r_docx.status_code == 200
    docx_doc = Document(io.BytesIO(r_docx.content))
    full_text = " ".join(p.text for p in docx_doc.paragraphs)
    # Check that event name, venue, lead name appear in the document
    assert unique_event_name in full_text
    assert "Darshil Prajapati" in full_text or any("Darshil Prajapati" in cell.text for t in docx_doc.tables for r in t.rows for cell in r.cells)

    # Verify PDF binary header
    r_pdf = client.get(f"/api/documents/download/{doc_id}/pdf")
    assert r_pdf.status_code == 200
    assert r_pdf.content.startswith(b"%PDF-")
    assert len(r_pdf.content) > 5000

